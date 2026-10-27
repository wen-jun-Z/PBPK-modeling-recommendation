from __future__ import annotations

import csv
import gzip
import json
import re
import shutil
import sqlite3
import subprocess
import zipfile
from collections import defaultdict
from hashlib import sha256
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import TreeBuilder
from xml.parsers.expat import errors as expat_errors
from xml.sax.handler import ContentHandler

from defusedxml import ElementTree as SafeET
from defusedxml import sax as SafeSAX

from pbpk_kb.parsers.dailymed_spl import DailyMedSPLParser
from pbpk_kb.parsers.emitter import ParseEmitter, clean_text, hash_text, json_safe, stable_id
from pbpk_kb.parsers.models import ParseAsset, ParseOutcome
from pbpk_kb.parsers.registry import ParserRegistry


TABLE_BLOCK_ROWS = 5_000
PUBLISHER_AVAILABILITY_MAX_BYTES = 16_384
PUBLISHER_AVAILABILITY_METADATA = {
    "publisher_metadata_kind": "data_availability",
    "scientific_evidence": False,
    "llm_eligible": False,
    "source_bytes_preserved": True,
    "recognition_version": "publisher-data-availability-v1",
}
JSON_FREE_TEXT_KEYS_BY_SOURCE = {
    "clinicaltrials": {"briefsummary", "detaileddescription", "eligibilitycriteria"},
    "openfda_drug_label": {
        "clinicalpharmacology",
        "pharmacokinetics",
        "druginteractions",
        "useinspecificpopulations",
        "pregnancy",
        "nursingmothers",
        "lactation",
        "pediatricuse",
        "geriatricuse",
        "renalimpairment",
        "hepaticimpairment",
        "precautions",
    },
    "pkdb": {"abstract", "text", "description", "comment"},
    "pubmed": {"abstract", "abstracttext"},
    "drugcentral": {"description", "pharmacology", "indication"},
}


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def element_text(element: Any | None) -> str:
    if element is None:
        return ""
    return clean_text(" ".join(element.itertext()))


def first_descendant(element: Any, name: str) -> Any | None:
    target = name.lower()
    for child in element.iter():
        if local_name(child.tag).lower() == target:
            return child
    return None


def descendants(element: Any, name: str) -> list[Any]:
    target = name.lower()
    return [child for child in element.iter() if local_name(child.tag).lower() == target]


def recognize_publisher_data_availability(payload: bytes) -> dict[str, str] | None:
    """Recognize one bounded publisher metadata schema, without repairing XML.

    Some PMC supplementary files use the literal root ``data:data`` without a
    namespace declaration. Namespace-disabled SAX can read that name as-is; it
    does not discard prefixes or recover malformed XML. DTDs and entities remain
    forbidden. Only the complete, whitelisted metadata shape below is accepted.
    Its numeric fields are publisher selection codes, never observations.
    """
    if len(payload) > PUBLISHER_AVAILABILITY_MAX_BYTES:
        return None
    try:
        lexical_text = payload.decode("utf-8-sig")
    except UnicodeDecodeError:
        return None
    # This compatibility schema is a UTF-8 envelope, not a recovery path for
    # unknown encodings. Keep lexical checks valid for the exact bytes parsed.
    if "\x00" in lexical_text:
        return None

    class MetadataHandler(ContentHandler):
        def __init__(self) -> None:
            super().__init__()
            self.tree = TreeBuilder()
            self.has_processing_instruction = False

        def startElement(self, name: str, attrs: Any) -> None:
            self.tree.start(name, dict(attrs))

        def endElement(self, name: str) -> None:
            self.tree.end(name)

        def characters(self, content: str) -> None:
            self.tree.data(content)

        def processingInstruction(self, target: str, data: str) -> None:
            self.has_processing_instruction = True

    handler = MetadataHandler()
    SafeSAX.parseString(payload, handler, forbid_dtd=True, forbid_entities=True, forbid_external=True)
    root = handler.tree.close()
    if handler.has_processing_instruction or "<!--" in lexical_text:
        return None

    def container(node: Any, tag: str, child_tags: list[str]) -> bool:
        return (
            node.tag == tag and not node.attrib and not (node.text or "").strip()
            and [child.tag for child in node] == child_tags
            and all(not (child.tail or "").strip() for child in node)
        )

    def leaf(node: Any, tag: str) -> bool:
        return node.tag == tag and not node.attrib and len(node) == 0

    if not container(root, "data:data", ["datasets"]):
        return None
    datasets = root[0]
    if not container(datasets, "datasets", ["author", "type", "dataset"]):
        return None
    author, kind, dataset = datasets
    if not leaf(author, "author") or not (author.text or "").strip() or not leaf(kind, "type"):
        return None
    if not container(dataset, "dataset", ["value", "reason", "comments"]):
        return None
    value, reason, comments = dataset
    if not all(leaf(node, tag) for node, tag in zip(dataset, ["value", "reason", "comments"])):
        return None
    if (
        (kind.text or "").strip() != "0" or (value.text or "").strip() != "1"
        or (reason.text or "").strip() != "Data will be made available on request"
        or (comments.text or "").strip()
    ):
        return None
    return {
        "availability_statement": "Data will be made available on request",
        "publisher_type": "0",
        "publisher_value": "1",
    }


class RawFormatParsers:
    def register(self, registry: ParserRegistry) -> None:
        registry.register("json", "json_parser", self.parse_json)
        registry.register("xml", "generic_xml_parser", self.parse_generic_xml)
        registry.register("xml", "dailymed_spl_parser", self.parse_dailymed_xml, source_id="dailymed")
        registry.register("xml", "pubmed_xml_parser", self.parse_pubmed_xml, source_id="pubmed")
        registry.register("xml", "pmc_jats_parser", self.parse_pmc_xml, source_id="pmc_oa")
        registry.register("html", "html_parser", self.parse_html)
        registry.register("pdf", "pdf_parser", self.parse_pdf)
        registry.register("table", "table_parser", self.parse_table)
        registry.register("document", "document_parser", self.parse_document)
        registry.register("presentation", "presentation_parser", self.parse_presentation)
        registry.register("text", "text_parser", self.parse_text)
        registry.register("image", "image_metadata_parser", self.parse_image)
        registry.register("media", "media_metadata_parser", self.parse_media)
        registry.register("database", "sqlite_parser", self.parse_database)
        registry.register("repo", "repository_parser", self.parse_repository)
        registry.register("repo", "osp_pksim_sqlite_parser", self.parse_database, source_id="osp_pbpk_model_library")
        registry.register("specialized_binary", "specialized_binary_parser", self.parse_specialized_binary)
        registry.register("unknown", "unknown_parser", self.parse_unknown)
        registry.register("filesystem_metadata", "filesystem_metadata_parser", self.parse_filesystem_metadata)

    def parse_filesystem_metadata(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        return ParseOutcome(status="metadata_only", metadata={
            "non_content_file_kind": asset.asset_metadata_json["non_content_file_kind"],
            "observation_extraction_blocked": True,
        })

    def parse_json(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        source_text = asset.physical_path.read_text(encoding="utf-8", errors="strict")
        tolerant = False
        warnings: list[str] = []
        try:
            payload = json.loads(source_text)
        except json.JSONDecodeError as strict_error:
            payload = json.loads(source_text, strict=False)
            tolerant = True
            warnings.append(f"Strict JSON parse failed; control-character fallback used: {strict_error}")
        emitter.record(
            asset,
            parser_name="json_parser",
            record_type="json_root",
            source_location=emitter.location(asset, "json_path=$"),
            record=payload,
            metadata={
                "top_level_type": type(payload).__name__,
                "request_key": asset.request_key,
                "tolerant_json_parse": tolerant,
            },
        )
        free_text_count = 0
        for record in iter_json_free_text(asset.source_id, payload):
            emitter.document(
                asset,
                parser_name="json_parser",
                content_kind="json_free_text",
                title=f"{asset.source_id} {record['key']}",
                text=record["text"],
                section_title=record["key"],
                source_location=emitter.location(asset, f"json_path={record['path']}"),
                metadata={"json_key": record["key"], "json_path": record["path"], "llm_eligible": True},
            )
            free_text_count += 1
        return ParseOutcome(
            warnings=warnings,
            metadata={"free_text_record_count": free_text_count, "tolerant_json_parse": tolerant},
        )

    def parse_dailymed_xml(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        payload = asset.physical_path.read_bytes()
        parsed = DailyMedSPLParser().parse(payload)
        root = SafeET.fromstring(payload)
        set_id_element = first_descendant(root, "setId")
        title = element_text(first_descendant(root, "title")) or display_name(asset)
        labeler = ""
        for organization in descendants(root, "representedOrganization"):
            labeler = element_text(first_descendant(organization, "name"))
            if labeler:
                break
        record = {
            "set_id": set_id_element.get("extension") if set_id_element is not None else None,
            "title": title,
            "labeler": labeler or None,
            **parsed,
        }
        emitter.record(
            asset,
            parser_name="dailymed_spl_parser",
            record_type="dailymed_spl",
            source_location=emitter.location(asset, "xml_path=/document"),
            record=record,
        )
        section_count = 0
        table_count = 0
        for section_index, section in enumerate(descendants(root, "section"), start=1):
            code_element = next((child for child in list(section) if local_name(child.tag) == "code"), None)
            section_code = code_element.get("code") if code_element is not None else None
            section_title = element_text(next((child for child in list(section) if local_name(child.tag) == "title"), None))
            text_element = next((child for child in list(section) if local_name(child.tag) == "text"), None)
            text = element_text(text_element)
            location = emitter.location(asset, f"xml_path=//section[{section_index}]")
            if text:
                emitter.document(
                    asset,
                    parser_name="dailymed_spl_parser",
                    content_kind="spl_section",
                    title=section_title or title,
                    section_title=section_title or None,
                    text=text,
                    source_location=location,
                    metadata={"section_code": section_code, "section_index": section_index, "llm_eligible": True},
                )
                section_count += 1
            if text_element is not None:
                for table_index, table in enumerate(descendants(text_element, "table"), start=1):
                    rows = xml_table_rows(table)
                    if rows:
                        columns = rows[0]
                        emitter.table(
                            asset,
                            parser_name="dailymed_spl_parser",
                            source_location=f"{location}#table={table_index}",
                            table_index=table_index,
                            columns=columns,
                            rows=rows[1:],
                            metadata={"section_code": section_code, "section_title": section_title, "llm_eligible": True},
                        )
                        table_count += 1
        return ParseOutcome(metadata={"section_count": section_count, "table_count": table_count})

    def parse_pubmed_xml(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        root = SafeET.parse(asset.physical_path).getroot()
        articles = descendants(root, "PubmedArticle")
        if not articles and local_name(root.tag) == "PubmedArticle":
            articles = [root]
        for article_index, article in enumerate(articles, start=1):
            pmid = element_text(first_descendant(article, "PMID"))
            title = element_text(first_descendant(article, "ArticleTitle"))
            journal = element_text(first_descendant(article, "Title"))
            abstract_parts = [element_text(node) for node in descendants(article, "AbstractText")]
            abstract = clean_text(" ".join(part for part in abstract_parts if part))
            authors = []
            for author in descendants(article, "Author"):
                name = clean_text(
                    " ".join(
                        part
                        for part in [
                            element_text(first_descendant(author, "ForeName")),
                            element_text(first_descendant(author, "LastName")),
                            element_text(first_descendant(author, "CollectiveName")),
                        ]
                        if part
                    )
                )
                if name:
                    authors.append(name)
            article_ids: dict[str, str] = {}
            for item in descendants(article, "ArticleId"):
                value = element_text(item)
                if value:
                    article_ids[str(item.get("IdType") or "unknown")] = value
            mesh_terms = [element_text(node) for node in descendants(article, "DescriptorName") if element_text(node)]
            publication_date = extract_pubmed_date(article)
            location = emitter.location(asset, f"xml_path=//PubmedArticle[{article_index}]")
            emitter.record(
                asset,
                parser_name="pubmed_xml_parser",
                record_type="pubmed_article",
                source_location=location,
                record={
                    "pmid": pmid or None,
                    "title": title or None,
                    "journal": journal or None,
                    "authors": authors,
                    "publication_date": publication_date,
                    "doi": article_ids.get("doi"),
                    "pmcid": article_ids.get("pmc"),
                    "mesh": mesh_terms,
                    "query_family": asset.request_key,
                },
            )
            if abstract:
                emitter.document(
                    asset,
                    parser_name="pubmed_xml_parser",
                    content_kind="abstract",
                    title=title or f"PMID {pmid}",
                    section_title="Abstract",
                    text=abstract,
                    source_location=f"{location}/Abstract",
                    metadata={"pmid": pmid, "doi": article_ids.get("doi"), "pmcid": article_ids.get("pmc"), "llm_eligible": True},
                )
        return ParseOutcome(status="parsed" if articles else "metadata_only", metadata={"article_count": len(articles)})

    def parse_pmc_xml(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        try:
            root = SafeET.parse(asset.physical_path).getroot()
        except SafeET.ParseError as error:
            if (
                error.code != expat_errors.codes[expat_errors.XML_ERROR_UNBOUND_PREFIX]
                or asset.source_id != "pmc_oa"
                or asset.physical_path.stat().st_size > PUBLISHER_AVAILABILITY_MAX_BYTES
            ):
                raise
            record = recognize_publisher_data_availability(asset.physical_path.read_bytes())
            if record is None:
                raise
            emitter.record(
                asset,
                parser_name="pmc_jats_parser",
                record_type="publisher_data_availability",
                source_location=emitter.location(asset, "xml_path=/data:data/datasets/dataset"),
                record=record,
                metadata=PUBLISHER_AVAILABILITY_METADATA,
            )
            return ParseOutcome(status="metadata_only", metadata=dict(PUBLISHER_AVAILABILITY_METADATA))
        articles = [node for node in root.iter() if local_name(node.tag) == "article"]
        if local_name(root.tag) == "article":
            articles = [root]
        if not articles:
            return self.parse_generic_xml(asset, emitter)
        total_tables = 0
        total_figures = 0
        total_paragraphs = 0
        for article_index, article in enumerate(articles, start=1):
            ids: dict[str, str] = {}
            for node in descendants(article, "article-id"):
                value = element_text(node)
                if value:
                    ids[str(node.get("pub-id-type") or "unknown")] = value
            pmcid = ids.get("pmc") or ids.get("pmcid")
            title = element_text(first_descendant(article, "article-title")) or f"PMC article {article_index}"
            article_location = emitter.location(asset, f"xml_path=//article[{article_index}]")
            emitter.record(
                asset,
                parser_name="pmc_jats_parser",
                record_type="pmc_article",
                source_location=article_location,
                record={"identifiers": ids, "pmcid": pmcid, "title": title},
            )
            for abstract_index, abstract_node in enumerate(descendants(article, "abstract"), start=1):
                abstract = element_text(abstract_node)
                if abstract:
                    emitter.document(
                        asset,
                        parser_name="pmc_jats_parser",
                        content_kind="abstract",
                        title=title,
                        section_title="Abstract",
                        text=abstract,
                        source_location=f"{article_location}/abstract[{abstract_index}]",
                        metadata={"pmcid": pmcid, "llm_eligible": True},
                    )
            parent_map = {child: parent for parent in article.iter() for child in list(parent)}
            sections = [node for node in article.iter() if local_name(node.tag) == "sec"]
            section_numbers = {node: index for index, node in enumerate(sections, start=1)}
            for paragraph_index, paragraph in enumerate(descendants(article, "p"), start=1):
                if not has_ancestor(paragraph, parent_map, "body"):
                    continue
                if has_any_ancestor(
                    paragraph,
                    parent_map,
                    {"ref-list", "table-wrap", "fig", "supplementary-material", "ack", "fn-group"},
                ):
                    continue
                text = element_text(paragraph)
                if not text:
                    continue
                section = nearest_ancestor(paragraph, parent_map, "sec")
                section_title = element_text(first_descendant(section, "title")) if section is not None else None
                section_index = section_numbers.get(section) if section is not None else None
                emitter.document(
                    asset,
                    parser_name="pmc_jats_parser",
                    content_kind="body_paragraph",
                    title=title,
                    section_title=section_title,
                    text=text,
                    source_location=f"{article_location}/sec[{section_index or 0}]/p[{paragraph_index}]",
                    metadata={"pmcid": pmcid, "section_index": section_index, "llm_eligible": True},
                )
                total_paragraphs += 1
            for table_index, table_wrap in enumerate(descendants(article, "table-wrap"), start=1):
                table_node = first_descendant(table_wrap, "table")
                rows = xml_table_rows(table_node) if table_node is not None else []
                caption = element_text(first_descendant(table_wrap, "caption"))
                if rows:
                    emitter.table(
                        asset,
                        parser_name="pmc_jats_parser",
                        source_location=f"{article_location}/table-wrap[{table_index}]",
                        table_index=table_index,
                        columns=rows[0],
                        rows=rows[1:],
                        metadata={"pmcid": pmcid, "caption": caption, "llm_eligible": True},
                    )
                    total_tables += 1
            for figure_index, figure in enumerate(descendants(article, "fig"), start=1):
                label = element_text(first_descendant(figure, "label"))
                caption = element_text(first_descendant(figure, "caption"))
                graphic = first_descendant(figure, "graphic")
                href = xml_href(graphic) if graphic is not None else None
                figure_location = f"{article_location}/fig[{figure_index}]"
                emitter.record(
                    asset,
                    parser_name="pmc_jats_parser",
                    record_type="figure",
                    source_location=figure_location,
                    record={"pmcid": pmcid, "label": label, "caption": caption, "href": href},
                )
                if href:
                    emitter.asset_link(
                        {
                            "run_id": asset.run_id,
                            "link_id": stable_id("asset_link", asset.raw_asset_id, figure_location, href),
                            "raw_asset_id": asset.raw_asset_id,
                            "source_id": asset.source_id,
                            "compound": asset.compound,
                            "checksum": asset.checksum,
                            "source_location": figure_location,
                            "link_type": "jats_graphic",
                            "target_href": href,
                            "target_raw_asset_id": None,
                            "metadata_json": {"pmcid": pmcid, "label": label, "caption": caption},
                        }
                    )
                total_figures += 1
            for supplement_index, supplement in enumerate(descendants(article, "supplementary-material"), start=1):
                href = xml_href(supplement)
                location = f"{article_location}/supplementary-material[{supplement_index}]"
                emitter.record(
                    asset,
                    parser_name="pmc_jats_parser",
                    record_type="supplement_link",
                    source_location=location,
                    record={"pmcid": pmcid, "href": href, "description": element_text(supplement)},
                )
                if href:
                    emitter.asset_link(
                        {
                            "run_id": asset.run_id,
                            "link_id": stable_id("asset_link", asset.raw_asset_id, location, href),
                            "raw_asset_id": asset.raw_asset_id,
                            "source_id": asset.source_id,
                            "compound": asset.compound,
                            "checksum": asset.checksum,
                            "source_location": location,
                            "link_type": "jats_supplement",
                            "target_href": href,
                            "target_raw_asset_id": None,
                            "metadata_json": {"pmcid": pmcid},
                        }
                    )
            for reference_index, reference in enumerate(descendants(article, "ref"), start=1):
                emitter.record(
                    asset,
                    parser_name="pmc_jats_parser",
                    record_type="reference",
                    source_location=f"{article_location}/ref[{reference_index}]",
                    record={"pmcid": pmcid, "text": element_text(reference)},
                )
        return ParseOutcome(
            metadata={
                "article_count": len(articles),
                "paragraph_count": total_paragraphs,
                "table_count": total_tables,
                "figure_count": total_figures,
            }
        )

    def parse_generic_xml(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        root = SafeET.parse(asset.physical_path).getroot()
        root_name = local_name(root.tag)
        title = element_text(first_descendant(root, "title")) or display_name(asset)
        emitter.record(
            asset,
            parser_name="generic_xml_parser",
            record_type="xml_root",
            source_location=emitter.location(asset, "xml_path=/"),
            record={"root_tag": root_name, "title": title, "attributes": dict(root.attrib)},
        )
        document_count = 0
        candidates = [node for node in root.iter() if local_name(node.tag).lower() in {"section", "sec", "abstract", "p"}]
        for index, node in enumerate(candidates, start=1):
            text = element_text(node)
            if len(text) < 40:
                continue
            tag = local_name(node.tag)
            emitter.document(
                asset,
                parser_name="generic_xml_parser",
                content_kind=f"xml_{tag.lower()}",
                title=title,
                section_title=element_text(first_descendant(node, "title")) or None,
                text=text,
                source_location=emitter.location(asset, f"xml_path=//{tag}[{index}]"),
                metadata={"root_tag": root_name, "llm_eligible": True},
            )
            document_count += 1
        if document_count == 0:
            text = element_text(root)
            if text:
                emitter.document(
                    asset,
                    parser_name="generic_xml_parser",
                    content_kind="xml_text",
                    title=title,
                    text=text,
                    source_location=emitter.location(asset, "xml_path=/"),
                    metadata={"root_tag": root_name, "llm_eligible": True},
                )
                document_count = 1
        return ParseOutcome(status="parsed" if document_count else "metadata_only", metadata={"root_tag": root_name})

    def parse_html(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        text = asset.physical_path.read_text(encoding="utf-8", errors="replace")
        parser = HTMLCollector()
        parser.feed(text)
        body = clean_text(" ".join(parser.text_parts))
        emitter.record(
            asset,
            parser_name="html_parser",
            record_type="html_metadata",
            source_location=emitter.location(asset, "css=html"),
            record={"title": parser.title, "headings": parser.headings, "links": parser.links},
        )
        if body:
            emitter.document(
                asset,
                parser_name="html_parser",
                content_kind="html_body",
                title=parser.title or display_name(asset),
                text=body,
                source_location=emitter.location(asset, "css=body"),
                metadata={"llm_eligible": True},
            )
        for index, rows in enumerate(parser.tables, start=1):
            if rows:
                emitter.table(
                    asset,
                    parser_name="html_parser",
                    source_location=emitter.location(asset, f"css=table:nth-of-type({index})"),
                    table_index=index,
                    columns=rows[0],
                    rows=rows[1:],
                    metadata={"llm_eligible": True},
                )
        return ParseOutcome(status="parsed" if body or parser.tables else "metadata_only")

    def parse_text(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        text = asset.physical_path.read_text(encoding="utf-8", errors="replace")
        document_id = emitter.document(
            asset,
            parser_name="text_parser",
            content_kind="text",
            title=display_name(asset),
            text=text,
            source_location=emitter.location(asset, "text=/"),
            metadata={"llm_eligible": asset.file_extension.lower() in {".txt", ".md"}},
        )
        return ParseOutcome(status="parsed" if document_id else "metadata_only")

    def parse_pdf(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        import fitz  # type: ignore

        document = fitz.open(str(asset.physical_path))
        if document.needs_pass:
            document.close()
            raise ValueError("Encrypted PDF requires a password")
        emitter.record(
            asset,
            parser_name="pdf_parser",
            record_type="pdf_metadata",
            source_location=emitter.location(asset, "pdf=/"),
            record={"page_count": len(document), "metadata": document.metadata},
        )
        text_pages = 0
        ocr_required_pages: list[int] = []
        table_count = 0
        for page_index in range(len(document)):
            page_number = page_index + 1
            page = document[page_index]
            raw_text = page.get_text("text")
            text = clean_text(raw_text)
            table_candidate, table_reasons = pdf_table_likelihood(raw_text, page.get_text("blocks"))
            page_location = emitter.location(asset, f"page={page_number}")
            if text:
                emitter.document(
                    asset,
                    parser_name="pdf_parser",
                    content_kind="pdf_page",
                    title=f"{display_name(asset)} page {page_number}",
                    text=text,
                    source_location=page_location,
                    metadata={
                        "page_number": page_number,
                        "llm_eligible": True,
                        "table_candidate": table_candidate,
                        "table_candidate_reasons": table_reasons,
                    },
                )
                text_pages += 1
            else:
                ocr_required_pages.append(page_number)
            find_tables = getattr(page, "find_tables", None)
            if table_candidate and callable(find_tables):
                try:
                    found = find_tables()
                    for page_table_index, table in enumerate(found.tables, start=1):
                        rows = table.extract() or []
                        if not rows:
                            continue
                        emitter.table(
                            asset,
                            parser_name="pdf_parser",
                            source_location=f"{page_location}#table={page_table_index}",
                            table_index=page_table_index,
                            page_number=page_number,
                            columns=list(rows[0]),
                            rows=[list(row) for row in rows[1:]],
                            metadata={"bounding_box": list(table.bbox), "llm_eligible": True},
                        )
                        table_count += 1
                except Exception:
                    pass
        document.close()
        status = "parsed" if text_pages or table_count else "metadata_only"
        warnings = [f"ocr_required_pages={len(ocr_required_pages)}"] if ocr_required_pages else []
        return ParseOutcome(
            status=status,
            warnings=warnings,
            metadata={
                "page_count": text_pages + len(ocr_required_pages),
                "text_page_count": text_pages,
                "ocr_required_pages": ocr_required_pages,
                "table_count": table_count,
            },
        )

    def parse_table(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        extension = asset.file_extension.lower()
        if extension in {".csv", ".tsv", ".tsv_gz"} or asset.physical_path.name.lower().endswith(".tsv.gz"):
            delimiter = "," if extension == ".csv" else "\t"
            if extension == ".tsv_gz" or asset.physical_path.name.lower().endswith(".tsv.gz"):
                with gzip.open(asset.physical_path, "rt", encoding="utf-8", errors="replace", newline="") as file:
                    count = self._emit_delimited(asset, emitter, csv.reader(file, delimiter=delimiter), "table_parser", None)
            else:
                with asset.physical_path.open("r", encoding="utf-8", errors="replace", newline="") as file:
                    count = self._emit_delimited(asset, emitter, csv.reader(file, delimiter=delimiter), "table_parser", None)
            return ParseOutcome(status="parsed" if count else "metadata_only", metadata={"table_block_count": count})
        if extension == ".xlsx":
            from openpyxl import load_workbook  # type: ignore

            with asset.physical_path.open("rb") as workbook_stream:
                workbook = load_workbook(workbook_stream, read_only=True, data_only=True)
                block_count = 0
                for sheet in workbook.worksheets:
                    block_count += self._emit_delimited(asset, emitter, sheet.iter_rows(values_only=True), "table_parser", sheet.title)
                workbook.close()
            return ParseOutcome(status="parsed" if block_count else "metadata_only", metadata={"table_block_count": block_count})
        if extension == ".xls":
            from pbpk_kb.parsers.source_integrity import verified_unavailable_ole

            unavailable = verified_unavailable_ole(asset)
            if unavailable is not None:
                return unavailable
            import xlrd  # type: ignore

            workbook = xlrd.open_workbook(asset.physical_path, on_demand=True)
            block_count = 0
            for sheet in workbook.sheets():
                rows = ([sheet.cell_value(row_index, column_index) for column_index in range(sheet.ncols)] for row_index in range(sheet.nrows))
                block_count += self._emit_delimited(asset, emitter, rows, "table_parser", sheet.name)
            workbook.release_resources()
            return ParseOutcome(status="parsed" if block_count else "metadata_only", metadata={"table_block_count": block_count})
        return ParseOutcome(status="unsupported", warnings=[f"No reliable parser for table extension {extension}"])

    def _emit_delimited(
        self,
        asset: ParseAsset,
        emitter: ParseEmitter,
        rows: Any,
        parser_name: str,
        sheet_name: str | None,
    ) -> int:
        iterator = iter(rows)
        try:
            columns = list(next(iterator))
        except StopIteration:
            return 0
        block: list[list[Any]] = []
        block_index = 0
        start_row = 2
        for row_number, row in enumerate(iterator, start=2):
            block.append(list(row))
            if len(block) >= TABLE_BLOCK_ROWS:
                block_index += 1
                emitter.table(
                    asset,
                    parser_name=parser_name,
                    source_location=emitter.location(asset, f"sheet={sheet_name or ''}#rows={start_row}-{row_number}"),
                    table_index=block_index,
                    sheet_name=sheet_name,
                    columns=columns,
                    rows=block,
                    metadata={"start_row": start_row, "end_row": row_number, "llm_eligible": True},
                )
                block = []
                start_row = row_number + 1
        if block:
            block_index += 1
            emitter.table(
                asset,
                parser_name=parser_name,
                source_location=emitter.location(asset, f"sheet={sheet_name or ''}#rows={start_row}-{start_row + len(block) - 1}"),
                table_index=block_index,
                sheet_name=sheet_name,
                columns=columns,
                rows=block,
                metadata={"start_row": start_row, "end_row": start_row + len(block) - 1, "llm_eligible": True},
            )
        return block_index

    def parse_document(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        extension = asset.file_extension.lower()
        if extension == ".docx":
            from docx import Document  # type: ignore

            document = Document(asset.physical_path)
            paragraph_count = 0
            for index, paragraph in enumerate(document.paragraphs, start=1):
                text = clean_text(paragraph.text)
                if not text:
                    continue
                emitter.document(
                    asset,
                    parser_name="document_parser",
                    content_kind="docx_paragraph",
                    title=display_name(asset),
                    section_title=paragraph.style.name if paragraph.style else None,
                    text=text,
                    source_location=emitter.location(asset, f"document.xml#paragraph={index}"),
                    metadata={"style": paragraph.style.name if paragraph.style else None, "llm_eligible": True},
                )
                paragraph_count += 1
            for table_index, table in enumerate(document.tables, start=1):
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                if rows:
                    emitter.table(
                        asset,
                        parser_name="document_parser",
                        source_location=emitter.location(asset, f"document.xml#table={table_index}"),
                        table_index=table_index,
                        columns=rows[0],
                        rows=rows[1:],
                        metadata={"llm_eligible": True},
                    )
            self._emit_ooxml_media_manifest(asset, emitter, "document_parser")
            return ParseOutcome(status="parsed" if paragraph_count or document.tables else "metadata_only")
        converter = shutil.which("antiword")
        command: list[str] | None = [converter, str(asset.physical_path)] if converter else None
        if command is None and shutil.which("textutil"):
            command = ["textutil", "-convert", "txt", "-stdout", str(asset.physical_path)]
        if command is None:
            return ParseOutcome(status="unsupported", warnings=["No antiword or textutil converter available"])
        completed = subprocess.run(command, capture_output=True, check=False, timeout=120)
        text = completed.stdout.decode("utf-8", errors="replace")
        if completed.returncode != 0 or not clean_text(text):
            return ParseOutcome(status="unsupported", warnings=[completed.stderr.decode("utf-8", errors="replace")[:1000]])
        emitter.document(
            asset,
            parser_name="document_parser",
            content_kind="legacy_document_text",
            title=display_name(asset),
            text=text,
            source_location=emitter.location(asset, "document=/"),
            metadata={"converter": Path(command[0]).name, "llm_eligible": True},
        )
        return ParseOutcome()

    def parse_presentation(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        if asset.file_extension.lower() != ".pptx":
            return ParseOutcome(status="unsupported", warnings=["No reliable legacy PPT parser available"])
        from pptx import Presentation  # type: ignore

        presentation = Presentation(asset.physical_path)
        slide_count = 0
        table_count = 0
        for slide_index, slide in enumerate(presentation.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    parts.extend(paragraph.text for paragraph in shape.text_frame.paragraphs if paragraph.text)
                if getattr(shape, "has_table", False):
                    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    if rows:
                        table_count += 1
                        emitter.table(
                            asset,
                            parser_name="presentation_parser",
                            source_location=emitter.location(asset, f"slide={slide_index}#table={table_count}"),
                            table_index=table_count,
                            columns=rows[0],
                            rows=rows[1:],
                            metadata={"slide_number": slide_index, "llm_eligible": True},
                        )
            notes = ""
            try:
                notes = clean_text(slide.notes_slide.notes_text_frame.text)
            except Exception:
                pass
            text = clean_text(" ".join(parts + ([notes] if notes else [])))
            if text:
                emitter.document(
                    asset,
                    parser_name="presentation_parser",
                    content_kind="pptx_slide",
                    title=f"{display_name(asset)} slide {slide_index}",
                    text=text,
                    source_location=emitter.location(asset, f"slide={slide_index}"),
                    metadata={"slide_number": slide_index, "notes_included": bool(notes), "llm_eligible": True},
                )
                slide_count += 1
        self._emit_ooxml_media_manifest(asset, emitter, "presentation_parser")
        return ParseOutcome(status="parsed" if slide_count or table_count else "metadata_only")

    def _emit_ooxml_media_manifest(self, asset: ParseAsset, emitter: ParseEmitter, parser_name: str) -> None:
        with zipfile.ZipFile(asset.physical_path) as archive:
            for info in archive.infolist():
                if "/media/" not in info.filename or info.is_dir():
                    continue
                with archive.open(info) as file:
                    digest = sha256()
                    total = 0
                    for chunk in iter(lambda: file.read(1024 * 1024), b""):
                        digest.update(chunk)
                        total += len(chunk)
                emitter.record(
                    asset,
                    parser_name=parser_name,
                    record_type="embedded_media",
                    source_location=emitter.location(asset, f"inner_path={info.filename}"),
                    record={"inner_path": info.filename, "byte_count": total, "sha256": digest.hexdigest()},
                )

    def parse_image(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        metadata: dict[str, Any] = {"byte_count": asset.file_size, "extension": asset.file_extension}
        warnings: list[str] = []
        try:
            from PIL import Image  # type: ignore

            with Image.open(asset.physical_path) as image:
                metadata.update(
                    {
                        "format": image.format,
                        "width": image.width,
                        "height": image.height,
                        "mode": image.mode,
                        "frame_count": getattr(image, "n_frames", 1),
                    }
                )
        except Exception as exc:
            warnings.append(f"Image metadata partially unavailable: {exc}")
        emitter.record(
            asset,
            parser_name="image_metadata_parser",
            record_type="image_metadata",
            source_location=emitter.location(asset, "image=/"),
            record=metadata,
            metadata={"ocr_or_vision_candidate": True},
        )
        return ParseOutcome(status="metadata_only", warnings=warnings, metadata=metadata)

    def parse_media(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        metadata: dict[str, Any] = {"byte_count": asset.file_size, "extension": asset.file_extension}
        warnings: list[str] = []
        try:
            from hachoir.metadata import extractMetadata  # type: ignore
            from hachoir.parser import createParser  # type: ignore

            parser = createParser(str(asset.physical_path))
            extracted = extractMetadata(parser) if parser else None
            if extracted:
                metadata["hachoir"] = {item.key: json_safe(item.value) for item in extracted}
        except Exception as exc:
            warnings.append(f"Media metadata partially unavailable: {exc}")
        emitter.record(
            asset,
            parser_name="media_metadata_parser",
            record_type="media_metadata",
            source_location=emitter.location(asset, "media=/"),
            record=metadata,
        )
        return ParseOutcome(status="metadata_only", warnings=warnings, metadata=metadata)

    def parse_database(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        uri = f"file:{asset.physical_path.as_posix()}?mode=ro&immutable=1"
        connection = sqlite3.connect(uri, uri=True)
        connection.execute("PRAGMA query_only = ON")
        table_names = [
            str(row[0])
            for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name")
        ]
        schema: list[dict[str, Any]] = []
        table_blocks = 0
        for table_index, table_name in enumerate(table_names, start=1):
            quoted = quote_identifier(table_name)
            columns_info = list(connection.execute(f"PRAGMA table_info({quoted})"))
            columns = [str(row[1]) for row in columns_info]
            row_count = int(connection.execute(f"SELECT COUNT(*) FROM {quoted}").fetchone()[0])
            schema.append(
                {
                    "table": table_name,
                    "row_count": row_count,
                    "columns": [
                        {"name": row[1], "type": row[2], "not_null": bool(row[3]), "primary_key": bool(row[5])}
                        for row in columns_info
                    ],
                }
            )
            cursor = connection.execute(f"SELECT * FROM {quoted}")
            start_row = 1
            block_index = 0
            while True:
                rows = cursor.fetchmany(TABLE_BLOCK_ROWS)
                if not rows:
                    break
                block_index += 1
                table_blocks += 1
                safe_rows = [[json_safe(cell) for cell in row] for row in rows]
                end_row = start_row + len(rows) - 1
                emitter.table(
                    asset,
                    parser_name="osp_pksim_sqlite_parser" if asset.source_id == "osp_pbpk_model_library" else "sqlite_parser",
                    source_location=emitter.location(asset, f"sqlite_table={table_name}#rows={start_row}-{end_row}"),
                    table_index=table_index * 100000 + block_index,
                    sheet_name=table_name,
                    columns=columns,
                    rows=safe_rows,
                    metadata={"table_row_count": row_count, "start_row": start_row, "end_row": end_row, "llm_eligible": False},
                )
                self._emit_database_text_fields(
                    asset,
                    emitter,
                    table_name,
                    columns,
                    safe_rows,
                    start_row,
                    "osp_pksim_sqlite_parser" if asset.source_id == "osp_pbpk_model_library" else "sqlite_parser",
                )
                start_row = end_row + 1
        emitter.record(
            asset,
            parser_name="osp_pksim_sqlite_parser" if asset.source_id == "osp_pbpk_model_library" else "sqlite_parser",
            record_type="sqlite_schema",
            source_location=emitter.location(asset, "sqlite_schema=/"),
            record={"sqlite_version": sqlite3.sqlite_version, "tables": schema},
        )
        connection.close()
        return ParseOutcome(metadata={"table_count": len(table_names), "table_block_count": table_blocks})

    def _emit_database_text_fields(
        self,
        asset: ParseAsset,
        emitter: ParseEmitter,
        table_name: str,
        columns: list[str],
        rows: list[list[Any]],
        start_row: int,
        parser_name: str,
    ) -> None:
        text_columns = {
            index
            for index, name in enumerate(columns)
            if re.search(r"(?:description|comment|abstract|text|notes?)", name, re.I)
        }
        for row_offset, row in enumerate(rows):
            for column_index in text_columns:
                if column_index >= len(row) or not isinstance(row[column_index], str) or len(clean_text(row[column_index])) < 40:
                    continue
                emitter.document(
                    asset,
                    parser_name=parser_name,
                    content_kind="sqlite_free_text",
                    title=f"{table_name}.{columns[column_index]}",
                    section_title=columns[column_index],
                    text=row[column_index],
                    source_location=emitter.location(
                        asset,
                        f"sqlite_table={table_name}#row={start_row + row_offset}#column={columns[column_index]}",
                    ),
                    metadata={"llm_eligible": True},
                )

    def parse_repository(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        path = asset.physical_path
        if path.is_file():
            with path.open("rb") as file:
                if file.read(16).startswith(b"SQLite format 3\x00"):
                    return self.parse_database(asset, emitter)
        if not path.is_dir():
            emitter.record(
                asset,
                parser_name="repository_parser",
                record_type="repository_file",
                source_location=emitter.location(asset, "file=/"),
                record={"file_name": display_name(asset), "byte_count": path.stat().st_size, "extension": asset.file_extension},
            )
            return ParseOutcome(status="metadata_only")
        files = []
        for item in path.rglob("*"):
            if item.is_file() and ".git" not in item.parts:
                files.append({"path": str(item.relative_to(path)), "byte_count": item.stat().st_size, "extension": item.suffix})
        emitter.record(
            asset,
            parser_name="repository_parser",
            record_type="repository_inventory",
            source_location=emitter.location(asset, "file=/"),
            record={"files": files},
        )
        return ParseOutcome(status="parsed" if files else "metadata_only", metadata={"file_count": len(files)})

    def parse_specialized_binary(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        with asset.physical_path.open("rb") as file:
            signature = file.read(32)
        emitter.record(
            asset,
            parser_name="specialized_binary_parser",
            record_type="specialized_binary_metadata",
            source_location=emitter.location(asset, "binary=/"),
            record={"extension": asset.file_extension, "byte_count": asset.file_size, "signature_hex": signature.hex()},
        )
        return ParseOutcome(status="unsupported", warnings=["No reliable read-only GraphPad Prism PZF parser available"])

    def parse_unknown(self, asset: ParseAsset, emitter: ParseEmitter) -> ParseOutcome:
        return ParseOutcome(status="unsupported", warnings=["No parser registered for unidentified binary content"])


class HTMLCollector(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.text_parts: list[str] = []
        self.links: list[str] = []
        self.headings: list[str] = []
        self.tables: list[list[list[str]]] = []
        self.title: str | None = None
        self._skip_depth = 0
        self._in_title = False
        self._in_heading = False
        self._table: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell_parts: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        lowered = tag.lower()
        if lowered in {"script", "style", "nav", "footer"}:
            self._skip_depth += 1
        if lowered == "title":
            self._in_title = True
        if lowered in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self._in_heading = True
        attrs_dict = {key.lower(): value for key, value in attrs}
        if lowered in {"a", "link"} and attrs_dict.get("href"):
            self.links.append(str(attrs_dict["href"]))
        if lowered == "table":
            self._table = []
        elif lowered == "tr" and self._table is not None:
            self._row = []
        elif lowered in {"td", "th"} and self._row is not None:
            self._cell_parts = []

    def handle_endtag(self, tag: str) -> None:
        lowered = tag.lower()
        if lowered in {"script", "style", "nav", "footer"} and self._skip_depth:
            self._skip_depth -= 1
        if lowered == "title":
            self._in_title = False
        if lowered in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self._in_heading = False
        if lowered in {"td", "th"} and self._cell_parts is not None and self._row is not None:
            self._row.append(clean_text(" ".join(self._cell_parts)))
            self._cell_parts = None
        elif lowered == "tr" and self._row is not None and self._table is not None:
            self._table.append(self._row)
            self._row = None
        elif lowered == "table" and self._table is not None:
            self.tables.append(self._table)
            self._table = None

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        text = clean_text(data)
        if not text:
            return
        self.text_parts.append(text)
        if self._in_title:
            self.title = text
        if self._in_heading:
            self.headings.append(text)
        if self._cell_parts is not None:
            self._cell_parts.append(text)


def iter_json_free_text(source_id: str, payload: Any) -> list[dict[str, str]]:
    allowed = JSON_FREE_TEXT_KEYS_BY_SOURCE.get(source_id, set())
    if not allowed:
        return []
    rows: list[dict[str, str]] = []

    def walk(value: Any, path: str, key: str | None = None) -> None:
        normalized_key = re.sub(r"[^a-z0-9]+", "", (key or "").lower())
        if isinstance(value, dict):
            for child_key, child_value in value.items():
                walk(child_value, f"{path}.{child_key}", str(child_key))
        elif isinstance(value, list):
            if normalized_key in allowed and all(isinstance(item, str) for item in value):
                add(normalized_key, path, "\n".join(value))
            else:
                for index, item in enumerate(value):
                    walk(item, f"{path}[{index}]", key)
        elif isinstance(value, str) and normalized_key in allowed:
            add(normalized_key, path, value)

    def add(key: str, path: str, value: str) -> None:
        text = clean_text(value)
        if len(text) >= 40:
            rows.append({"key": key, "path": path, "text": text})

    walk(payload, "$")
    seen: set[tuple[str, str]] = set()
    unique = []
    for row in rows:
        key = (row["path"], hash_text(row["text"]))
        if key not in seen:
            seen.add(key)
            unique.append(row)
    return unique


def extract_pubmed_date(article: Any) -> str | None:
    date_node = first_descendant(article, "PubDate")
    if date_node is None:
        date_node = first_descendant(article, "ArticleDate")
    if date_node is None:
        return None
    year = element_text(first_descendant(date_node, "Year"))
    month = element_text(first_descendant(date_node, "Month"))
    day = element_text(first_descendant(date_node, "Day"))
    return "-".join(part for part in [year, month, day] if part) or None


def nearest_ancestor(element: Any, parent_map: dict[Any, Any], tag: str) -> Any | None:
    current = parent_map.get(element)
    while current is not None:
        if local_name(current.tag) == tag:
            return current
        current = parent_map.get(current)
    return None


def has_ancestor(element: Any, parent_map: dict[Any, Any], tag: str) -> bool:
    return nearest_ancestor(element, parent_map, tag) is not None


def has_any_ancestor(element: Any, parent_map: dict[Any, Any], tags: set[str]) -> bool:
    current = parent_map.get(element)
    while current is not None:
        if local_name(current.tag) in tags:
            return True
        current = parent_map.get(current)
    return False


def xml_table_rows(table: Any | None) -> list[list[str]]:
    if table is None:
        return []
    rows: list[list[str]] = []
    for row in descendants(table, "tr"):
        cells = [element_text(cell) for cell in list(row) if local_name(cell.tag).lower() in {"td", "th"}]
        if cells:
            rows.append(cells)
    return rows


def xml_href(element: Any | None) -> str | None:
    if element is None:
        return None
    for key, value in element.attrib.items():
        if local_name(key).lower() == "href":
            return str(value)
    return None


def quote_identifier(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


def display_name(asset: ParseAsset) -> str:
    return Path(asset.inner_path).name if asset.inner_path else Path(asset.storage_path).name


def pdf_table_likelihood(text: str, blocks: list[Any]) -> tuple[bool, list[str]]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    reasons: list[str] = []
    numeric_lines = sum(1 for line in lines if len(re.findall(r"[-+]?\d+(?:\.\d+)?", line)) >= 2)
    delimited_lines = sum(1 for line in lines if "\t" in line or re.search(r"\S\s{3,}\S", line))
    pk_header = any(
        re.search(r"\b(?:Cmax|AUC|Tmax|half[- ]life|clearance|dose|concentration|parameter|unit)\b", line, re.I)
        for line in lines
    )
    if numeric_lines >= 3:
        reasons.append("numeric_line_density")
    if delimited_lines >= 2:
        reasons.append("aligned_text_columns")
    if pk_header and numeric_lines >= 2:
        reasons.append("pk_header_with_numeric_rows")
    if len(blocks) >= 12 and numeric_lines >= 2:
        reasons.append("dense_layout_blocks")
    return bool(reasons), reasons
