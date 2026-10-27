#!/usr/bin/env python3
"""Freeze 100 structure-isolated, zero-curve PBPK construction cases.

HT-PBK supplies a drug/IV source index and predicted physicochemical inputs.
PKSmart supplies training rows and separately held-out aggregate human PK
references. These are NOT 100 matched clinical concentration-time studies.
"""
from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
import math
from pathlib import Path
import re
import sys
import unicodedata

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))

import openpyxl
from rdkit import Chem, rdBase
from rdkit.Chem import Descriptors, rdMolDescriptors
from rdkit.Chem.MolStandardize import rdMolStandardize
from pbpk_kb.blueprints.pk_parameter_learning import load_parameter_dataset, normalize_structure
from pbpk_kb.blueprint_learning.ranged_parameter_prior import fit_ranged_prior
from pbpk_kb.blueprint_policy.purpose_design import plan_design

VERSION = 'purpose-hundred-cohort-1'
DEFAULT = ROOT / 'experiments/blueprint_design/purpose_hundred_20260930'
WHEEL = ROOT / 'artifacts/pk_predictors/pksmart/pksmart-3.0.1-py3-none-any.whl'
ENDPOINTS = ['plasma_concentration_time', 'auc', 'cmax', 'tmax', 'c_last',
             'terminal_half_life', 'auc_inf', 'clearance', 'vz', 'mrt', 'vss']
REFERENCES = {'clearance': ('human_CL_mL_min_kg', 'l/h', 70 * 60 / 1000),
              'vss': ('human_VDss_L_kg', 'l', 70),
              'terminal_half_life': ('human_thalf', 'h', 1),
              'mrt': ('human_mrt', 'h', 1)}


def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda: f.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def ref(path, location=None):
    p = Path(path).resolve(strict=True)
    if not p.is_relative_to(ROOT):
        raise ValueError('Source escaped project')
    out = {'path': str(p.relative_to(ROOT)), 'sha256': sha(p)}
    if location is not None:
        out['location'] = location
    return out


def save_new(path, obj):
    p = Path(path).resolve()
    if not p.is_relative_to(DEFAULT):
        raise ValueError('Preparation writes only inside this dedicated batch')
    data = (json.dumps(obj, sort_keys=True, indent=2, ensure_ascii=False, allow_nan=False) + '\n').encode()
    if p.exists():
        if p.read_bytes() != data:
            raise FileExistsError('Preserving different existing artifact: ' + str(p))
    else:
        p.parent.mkdir(parents=True, exist_ok=True)
        with p.open('xb') as f:
            f.write(data)
    return ref(p)


def normalized_name(value):
    return unicodedata.normalize('NFKC', str(value)).strip().casefold()


def identifier(identity):
    return 'structure:' + hashlib.sha256(identity.encode()).hexdigest()[:20]


def stereo_parent(smiles):
    """Stricter reference matching than the conservative RF exclusion key."""
    with rdBase.BlockLogs():
        mol = Chem.MolFromSmiles(smiles)
        if mol is None:
            raise ValueError('Invalid stereochemical input')
        mol = rdMolStandardize.Uncharger().uncharge(rdMolStandardize.FragmentParent(mol))
        # Do not canonicalize tautomers here: that can erase stereochemistry at
        # enolizable centers. Unequal but possibly equivalent source structures
        # stay unresolved, rather than silently being declared clinically equal.
        return Chem.MolToSmiles(mol, isomericSmiles=True)


def freeze_selection(candidates, expected=100):
    """Deduplicate then select by identity only, never by labels or predictions."""
    unique = {}
    for row in sorted(candidates, key=lambda x: (x['identity_smiles'], x['ht_row'])):
        key = row['identity_smiles']
        if key not in unique:
            unique[key] = dict(row, aliases=[row['drug_name']])
        elif row['drug_name'] not in unique[key]['aliases']:
            unique[key]['aliases'].append(row['drug_name'])
    ranked = sorted(unique.values(), key=lambda x: hashlib.sha256(
        ('20260930|' + x['identity_smiles']).encode()).hexdigest())
    if len(ranked) < expected:
        raise ValueError(f'Only {len(ranked)} unique eligible identities; require {expected}')
    return ranked[:expected], ranked[expected:]


def collect_candidates():
    sources = DEFAULT / 'sources'
    chemical_path = sources / '204_2024_3764_MOESM2_ESM.xlsx'
    index_path = sources / '204_2024_3764_MOESM3_ESM.xlsx'
    dataset = load_parameter_dataset(WHEEL)
    # First standardized occurrence is fixed without using endpoint values.
    label_rows = {}
    for row in dataset.rows:
        if row['identity_smiles']:
            label_rows.setdefault(row['identity_smiles'], row)
    book = openpyxl.load_workbook(index_path, read_only=True, data_only=True)
    sheet = book.active
    values = sheet.iter_rows(values_only=True)
    heads = next(values)
    studies = {}
    for excel_row, cells in enumerate(values, 2):
        row = dict(zip(heads, cells))
        if normalized_name(row.get('route')) in {'iv', 'intravenous'}:
            name = normalized_name(row['compoundName'])
            studies.setdefault(name, []).append({'excel_row': excel_row,
                'database': row.get('database'), 'databaseID': row.get('databaseID'),
                'referenceName': row.get('referenceName'), 'referenceLink': row.get('referenceLink'),
                'pubmedID': row.get('pubmedID')})
    book.close()
    book = openpyxl.load_workbook(chemical_path, read_only=True, data_only=True)
    values = book['Preds'].iter_rows(values_only=True)
    next(values)
    heads = next(values)
    candidates, exclusions = [], []
    for excel_row, cells in enumerate(values, 3):
        row = dict(zip(heads, cells))
        if not row.get('compoundName'):
            continue
        name = normalized_name(row['compoundName'])
        if name not in studies:
            exclusions.append({'ht_row': excel_row, 'drug_name': name, 'reason': 'no_IV_index_record'})
            continue
        try:
            identity = normalize_structure(row['SMILES'])
        except (ValueError, TypeError) as exc:
            exclusions.append({'ht_row': excel_row, 'drug_name': name,
                               'reason': 'structure_normalization_failed', 'detail': str(exc)})
            continue
        if identity['identity_smiles'] not in label_rows:
            exclusions.append({'ht_row': excel_row, 'drug_name': name, 'reason': 'not_in_PKSmart_identity_index'})
            continue
        pka = []
        for key, kind in [('pKa-a_Pred1', 'Acid'), ('pKa-b_Pred1', 'Base')]:
            value = row.get(key)
            if isinstance(value, (int, float)) and math.isfinite(value):
                pka.append({'Type': kind, 'Pka': float(value)})
            elif value is not None:
                raise ValueError(f'Unexpected pKa type at row {excel_row}: {value!r}')
        if not pka:
            exclusions.append({'ht_row': excel_row, 'drug_name': name, 'reason': 'no_predicted_pKa_available'})
            continue
        mol = Chem.MolFromSmiles(identity['identity_smiles'])
        if mol is None:
            raise ValueError('Normalized molecular identity became invalid')
        human_name = re.sub(r'[^A-Za-z0-9 _-]', '', unicodedata.normalize('NFKD', name).encode('ascii', 'ignore').decode())
        if not human_name or not human_name[0].isalpha():
            human_name = 'Drug ' + human_name
        pk_row = label_rows[identity['identity_smiles']]
        ht_stereo, pk_stereo = stereo_parent(row['SMILES']), stereo_parent(pk_row['smiles_r'])
        candidates.append({**identity, 'compound_id': identifier(identity['identity_smiles']),
            'drug_name': human_name[:80], 'source_drug_name': name, 'ht_row': excel_row,
            'source_cas': row.get('compoundCas'),
            'ht_stereo_parent': ht_stereo, 'pk_stereo_parent': pk_stereo,
            'primary_reference_eligible': ht_stereo == pk_stereo,
            'reference_identity_status': 'exact_standardized_stereo_match' if ht_stereo == pk_stereo else 'stereo_agreement_unresolved',
            'pk_row_index': label_rows[identity['identity_smiles']]['index'],
            'pk_csv_line': label_rows[identity['identity_smiles']]['csv_line'],
            'input_smiles': row['SMILES'], 'pka_types': pka,
            'pka_sources': [str(row.get('pKa-a_Source1')), str(row.get('pKa-b_Source1'))],
            'molecular_weight': float(Descriptors.MolWt(mol)),
            'molecular_formula': rdMolDescriptors.CalcMolFormula(mol),
            'logp': float(Descriptors.MolLogP(mol)), 'IV_index_records': studies[name]})
    book.close()
    return candidates, exclusions, dataset, [chemical_path, index_path, WHEEL]


def evaluation_record(row, label, source):
    endpoints = {}
    for key, (column, unit, multiplier) in REFERENCES.items():
        log_value = label['labels_log10'].get(column)
        raw_value = None if log_value is None else 10 ** log_value * multiplier
        value = raw_value if row['primary_reference_eligible'] else None
        endpoints[key] = {'value': value, 'unit': unit, 'source_column': column,
            'connectivity_matched_value_not_scored': raw_value if not row['primary_reference_eligible'] else None,
            'reference_identity_status': row['reference_identity_status'],
            'exclusion_reason': None if row['primary_reference_eligible'] else 'stereo_agreement_unresolved',
            'source_csv_line': label['csv_line'], 'stored_scale': 'log10',
            'conversion_multiplier': multiplier,
            'basis': 'published_aggregate_human_IV_PK_reference_not_matched_study',
            'source': source}
    return {'compound_id': row['compound_id'], 'identity_smiles': row['identity_smiles'],
        'evaluation_scope': 'aggregated_human_IV_PK_parameters', 'weight_kg': 70.,
        'label_role': 'evaluation_only', 'endpoints': endpoints, 'clinical_curve': None,
        'reference_identity_status': row['reference_identity_status'],
        'source': source, 'study_id': None,
        'limitations': [
            'PKSmart CSV has no original study identifiers, dosing, demographic or matrix records.',
            'Per-kg references are converted to the hypothetical 70 kg adult, not a matched patient.',
            'MRT and half-life may differ with source sampling, infusion and NCA conventions.',
            'References do not validate a hypothetical dose curve or all eleven endpoints.',
            'First normalized row selected without using endpoint availability or value; duplicates are not additional drugs.',
            'Unresolved stereochemical matching is retained for construction but excluded from primary accuracy scoring.']}


def prepare(n_jobs=4):
    candidates, excluded, dataset, paths = collect_candidates()
    cohort, unselected = freeze_selection(candidates)
    old_path = ROOT / 'experiments/blueprint_design/optimization_20260929/priors/parameter_split.json'
    old = json.loads(old_path.read_text())['external_exclusions']
    old_identities = {r['identity_smiles'] for r in old.values()}
    for row in cohort:
        row['cohort_role'] = 'previously_viewed' if row['identity_smiles'] in old_identities else 'development_extension'
        row['case_id'] = 'drug_' + row['compound_id'].split(':')[1]
    sources = {str(p.relative_to(ROOT)): sha(p) for p in paths + [old_path, Path(__file__)]}
    registry = {'version': VERSION, 'selection_rule': 'first HT row per normalized identity, then SHA256(20260930|identity) ascending, first 100',
        'selection_uses_labels_or_predictions': False, 'planned_compounds': 100,
        'candidates_with_aliases': len(candidates), 'unique_eligible': len(cohort) + len(unselected),
        'cases': cohort, 'exclusions': excluded,
        'eligible_unselected_ids': [r['compound_id'] for r in unselected],
        'source_pins': sources,
        'scientific_scope': 'construction and held-out aggregate parameter benchmark; not 100 clinical curves',
        'phase': 'development_diagnostic_not_confirmatory',
        'primary_reference_eligible': sum(r['primary_reference_eligible'] for r in cohort),
        'identity_semantics': '100 distinct normalized connectivities. HT R/S aliases sharing one connectivity are source-name variants, not asserted stereochemical synonyms.',
        'physical_input_notes': 'RDKit normalized parent MW/formula/logP; HT-PBK ChemAxon predicted pKa. Predictor training overlap unknown.'}
    save_new(DEFAULT / 'cohort_registry.json', registry)
    exclusions = {r['compound_id']: r['identity_smiles'] for r in cohort}
    prior = fit_ranged_prior(WHEEL, exclude_smiles=exclusions,
        output_dir=DEFAULT / 'priors', project_root=ROOT, n_jobs=n_jobs)
    assumptions = {'weight_kg': 70., 'age_years': 30., 'height_cm': 175., 'sex': 'MALE',
        'population': 'European_ICRP_2002', 'dose_mg': 70., 'infusion_minutes': 30.,
        'window_h': [0., 48.], 'solubility': 1., 'solubility_ph': 7.4,
        'binding_partner': 'Unknown', 'elimination_scope': 'hepatic_only',
        'description': 'Hypothetical already-dissolved single IV engineering scenario. Lumped systemic clearance, not validated hepatic mechanism.',
        'endpoint_limits': {e: {'metric': 'absolute' if e == 'tmax' else 'fold',
                               'max_change': 1. if e == 'tmax' else 2.} for e in ENDPOINTS},
        'threshold_interpretation': 'Predeclared stress-screening tolerance, not clinical safety or certification.'}
    assumption_ref = save_new(DEFAULT / 'inputs' / 'assumptions.json', assumptions)
    physical_ref = save_new(DEFAULT / 'inputs' / 'physical_inputs.json', {
        'rows': {r['compound_id']: {k: r[k] for k in ('ht_row', 'input_smiles', 'identity_smiles',
            'molecular_weight', 'molecular_formula', 'logp', 'pka_types', 'pka_sources')} for r in cohort},
        'source': ref(paths[0]), 'rdkit_version': rdBase.rdkitVersion,
        'contains_clinical_evaluation_labels': False,
        'ionization_scope': 'At most one reported predicted acid and base pKa. Missing categories do not establish measured neutrality.'})
    all_sources = dict(sources)
    all_sources.update({str((DEFAULT / 'cohort_registry.json').relative_to(ROOT)): sha(DEFAULT / 'cohort_registry.json')})
    for item in [assumption_ref, physical_ref, *prior['artifacts'].values()]:
        all_sources[item['path']] = item['sha256']
    rows_by_index = {r['index']: r for r in dataset.rows}
    cases = []
    for row in cohort:
        cid, case_id = row['compound_id'], row['case_id']
        forecast = prior['predictions'][cid]
        ev = evaluation_record(row, rows_by_index[row['pk_row_index']], ref(WHEEL))
        evaluation_ref = save_new(DEFAULT / 'evaluation' / (case_id + '.json'), ev)
        evidence = []
        def add(parameter, value, unit, basis, source, notes='', prediction=None, context=None):
            entry = {'parameter': parameter, 'compound_id': cid, 'value': value, 'unit': unit,
                'basis': basis, 'source': dict(source, location=f'{cid}/{parameter}'),
                'notes': notes, 'context': context or {}}
            if prediction:
                entry['prediction_method'] = prediction
            evidence.append(entry)
        descriptor = {'name': 'RDKit MolLogP', 'version': rdBase.rdkitVersion,
            'training_overlap': 'unknown', 'applicability': 'Structure descriptor only; no held-out clinical endpoint inputs'}
        add('molecular_weight', row['molecular_weight'], 'g/mol', 'derived', physical_ref)
        add('molecular_formula', row['molecular_formula'], 'text', 'derived', physical_ref)
        add('logp', row['logp'], 'Log Units', 'predicted', physical_ref, prediction=descriptor)
        add('pka_types', row['pka_types'], '1', 'predicted', physical_ref,
            prediction={'name': 'ChemAxon via HT-PBK', 'version': '2024 published supplementary predictions',
                'training_overlap': 'unknown', 'applicability': 'Published chemical pKa prediction; no claim of independently measured pKa'})
        for key, unit in [('solubility', 'mg/ml'), ('solubility_ph', '1'), ('binding_partner', 'text'), ('elimination_scope', 'text')]:
            add(key, assumptions[key], unit, 'assumption', assumption_ref, notes=assumptions['description'])
        context = {'species': 'human', 'matrix': 'plasma', 'route': 'intravenous',
                   'population_condition': 'reference_adult', 'clearance_type': 'systemic'}
        predictor = {'name': 'random_forest', 'version': prior['audit']['version'], 'training_overlap': False,
            'applicability': {'human_IV_parameter_model': True, 'target_structure_and_scaffold_excluded': True}}
        for parameter, unit in [('total_clearance', 'ml/min/kg'), ('fraction_unbound', 'fraction')]:
            add(parameter, forecast['predict_nominal'][parameter], unit, 'predicted', prior['artifacts']['predictions'],
                context=context, prediction=predictor)
        ledger = save_new(DEFAULT / 'inputs' / case_id / 'evidence.json', {'evidence': evidence})
        request = {'entry_reason': 'model_missing', 'compound_id': cid, 'compound_name': row['drug_name'],
            'human': {k: assumptions[k] for k in ('weight_kg', 'age_years', 'height_cm', 'sex', 'population')} | {'seed': 17},
            'dosing': {'route': 'intravenous', 'dose_mg': assumptions['dose_mg'], 'start_h': 0., 'infusion_minutes': 30.},
            'target': {'endpoints': ENDPOINTS, 'primary_endpoint': 'plasma_concentration_time', 'window_start_h': 0., 'window_end_h': 48.},
            'partition_method': 'PKSimStandard', 'evidence_ledger': dict(ledger, location='Full frozen construction evidence'),
            'evidence': evidence, 'forbidden_sources': [dict(evaluation_ref, location='Held-out aggregate human PK outcomes; scoring after all predictions locked')],
            'assumptions': [assumptions['description'],
                'No clinical concentration points or compound-specific fitted parameters enter construction.',
                'No Vss, half-life or MRT reference is used to set distribution or simulate the tail.',
                'RF train and range calibration exclude all 100 compound identities and scaffolds.',
                'Predicted pKa categories are incomplete chemical inputs, not clinical measurements.',
                'Missing organ clearance split is lumped as systemic clearance; hepatic_only is a construction proxy.',
                'No complete PBPK model is loaded, even if one exists for a selected drug.',
                'Sensitivity ranges are empirical marginal stress bounds; simultaneous corners are not a joint confidence interval.']}
        ranges = [{'parameter': p, 'lower': v['lower'], 'upper': v['upper'], 'unit': v['unit'],
            'basis': 'training_validation', 'source': dict(prior['artifacts']['predictions'], location=f'{cid}/ranges/{p}')}
            for p, v in forecast['ranges'].items() if not v['degenerate']]
        document = {'request': request, 'fallback_evidence': [], 'ranges': ranges,
                    'endpoint_limits': assumptions['endpoint_limits']}
        plan = plan_design(request, [], ranges, endpoint_limits=assumptions['endpoint_limits'])
        if plan['status'] != 'ready':
            raise ValueError(f'Invalid prepared input {case_id}: {plan}')
        inp = save_new(DEFAULT / 'inputs' / case_id / 'request.json', document)
        cases.append({k: row[k] for k in ('case_id', 'compound_id', 'identity_smiles', 'drug_name', 'cohort_role')} |
            {'evaluation_scope': ev['evaluation_scope'], 'input': inp, 'evaluation': evaluation_ref,
             'evaluation_sources': [ref(WHEEL)]})
        for item in (evaluation_ref, ledger, inp):
            all_sources[item['path']] = item['sha256']
    manifest = {'version': VERSION, 'cases': cases, 'source_pins': all_sources,
        'scientific_scope': registry['scientific_scope'], 'phase': registry['phase'],
        'declared_comparisons': ['purpose-conditioned endpoint sensitivity', 'FixedAuto same nominal predictions', 'input-only range-width ranking'],
        'clinical_curve_count': 0, 'independent_study_validation': False,
        'previously_viewed_compounds': sum(r['cohort_role'] == 'previously_viewed' for r in cohort),
        'selection_uses_labels_or_predictions': False,
        'ref_training_is_scaffold_disjoint': True,
        'reference_semantics': 'Four aggregate human IV PK parameter labels, not a single matched study or 100 clinical curves.'}
    save_new(DEFAULT / 'cohort.json', manifest)
    return {'manifest': ref(DEFAULT / 'cohort.json'), 'cases': len(cases),
        'unique_eligible': registry['unique_eligible'], 'viewed': manifest['previously_viewed_compounds'],
        'reference_counts': {e: sum(r['primary_reference_eligible'] and rows_by_index[r['pk_row_index']]['labels_log10'].get(col) is not None for r in cohort)
                             for e, (col, _, _) in REFERENCES.items()},
        'clinical_curve_count': 0, 'planned_range_scenarios': 900}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--n-jobs', type=int, default=4)
    args = parser.parse_args()
    print(json.dumps(prepare(args.n_jobs), indent=2))
