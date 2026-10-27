#!/usr/bin/env python3
"""Build the English blueprint-method deck combining diagrams and brief explanations.

The PPTX and vector PDF share one layout. Equations are embedded as high-resolution
images in PPTX and vectors in PDF. All other text, boxes and arrows remain editable.
No scientific inputs, fitted models, paper assets or experimental outputs are changed.
"""
from __future__ import annotations

import io
import math
import os
import re
from pathlib import Path
import unicodedata
import zipfile

ROOT = Path(__file__).resolve().parents[1]
os.environ.setdefault("MPLCONFIGDIR", str(ROOT / "runtime/matplotlib-cache"))
os.environ.setdefault("PYTHONDONTWRITEBYTECODE", "1")

import matplotlib
matplotlib.use("Agg")
from matplotlib.figure import Figure
from matplotlib.backends.backend_agg import FigureCanvasAgg
import pymupdf as fitz
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

OUT = ROOT / "results/presentations"
STEM = 'PBPK_blueprint_principles_and_workflow'
FONT_FILE = ROOT / "runtime/fonts/NotoSansCJKsc-Regular.otf"
FONT_NAME = "Noto Sans CJK SC"
FONT = fitz.Font(fontfile=str(FONT_FILE))
BOLD_FILE = ROOT / "runtime/fonts/noto_serif_cjk/NotoSerifCJKsc-Bold.otf"
BOLD_NAME = "Noto Serif CJK SC"
BOLD_FONT = fitz.Font(fontfile=str(BOLD_FILE))
W, H = 960, 540
EQUATION_FONT_SIZE = 24
TEAL, INK, MUTED = "255E72", "26353F", "657780"
LINE, PALE, WHITE = "91AEB7", "F2F6F7", "FFFFFF"
GREEN, GREEN_PALE = "427964", "F0F6F2"
GOLD, GOLD_PALE = "A87949", "FAF5EC"
PURPLE, PURPLE_PALE = "806C9D", "F5F2F8"


def rgb(s):
    return tuple(int(s[i:i+2], 16) / 255 for i in (0, 2, 4))


def wrap(text, width, size):
    lines = []
    for paragraph in str(text).split("\n"):
        tokens = re.findall(r"[A-Za-z0-9]+(?:[./_-][A-Za-z0-9]+)*|\s+|.", paragraph)
        grouped = []
        for token in tokens:
            if token in "，。；：！？、）】》〉”’％%,.;:!?)]}" and grouped:
                grouped[-1] += token
            else:
                grouped.append(token)
        line = ""
        for token in grouped:
            if line and FONT.text_length(line + token, fontsize=size) > width - 1:
                lines.append(line.rstrip())
                line = token.lstrip()
            else:
                line += token
        lines.append(line)
    return lines


def fit_size(text, width, height, size, *, line=1.2, minimum=14):
    """Fit translated prose without changing diagram geometry or equations."""
    for candidate in range(int(size), minimum-1, -1):
        rows = wrap(text, width, candidate)
        if (len(rows)*candidate*line+3 <= height
                and all(FONT.text_length(row, fontsize=candidate) <= width for row in rows)):
            return candidate
    raise ValueError('Shorten text to fit the layout: '+str(text))


class Slide:
    def __init__(self, title, subtitle, section, notes):
        self.ops = []
        self.notes = notes
        self.text(42, 24, 800, "PBPK BLUEPRINT   /   " + section, 10, TEAL, bold=True)
        self.text(42, 55, 876, title, fit_size(title, 876, 44, 29, line=1.4, minimum=22), TEAL)
        self.text(42, 103, 876, subtitle, fit_size(subtitle, 876, 23, 14, line=1.4, minimum=11), MUTED)
        self.line([(42, 494), (918, 494)], "DCE5E8", .8)
        self.text(42, 505, 805, 'Blueprint design · Inputs—computation—outputs · 2026.10.13', 9, MUTED)

    def text(self, x, y, w, value, size=18, color=INK, bold=False, align="left", line=1.4):
        rows = wrap(value, w, size)
        h = len(rows) * size * line + 3
        if x < 0 or y < 0 or x + w > W + .1 or y + h > H:
            raise ValueError(f"Out-of-slide text: {value}")
        for c in str(value):
            if c not in "\n\t" and not FONT.has_glyph(ord(c)):
                raise ValueError(f"Missing glyph {c!r} in {value!r}")
        self.ops.append(dict(kind="text", x=x, y=y, w=w, h=h, rows=rows,
                             size=size, color=color, bold=bold, align=align, spacing=line))
        return h

    def rect(self, x, y, w, h, fill=PALE, stroke=LINE, width=1, rounded=True):
        self.ops.append(dict(kind="rect", x=x, y=y, w=w, h=h, fill=fill,
                             stroke=stroke, width=width, rounded=rounded))

    def line(self, points, color=TEAL, width=1.4, dash=False, arrow=False, both=False):
        self.ops.append(dict(kind="line", points=points, color=color, width=width,
                             dash=dash, arrow=arrow, both=both))

    def eq(self, x, y, w, h, latex, size=EQUATION_FONT_SIZE, color=INK, *, role="formula"):
        self.ops.append(dict(kind="equation", x=x, y=y, w=w, h=h,
                             latex=latex, size=size, color=color, role=role))

    def ellipse(self, x, y, w, h=None, fill=PALE, stroke=TEAL, width=1.3):
        self.ops.append(dict(kind="ellipse", x=x, y=y, w=w, h=h or w,
                             fill=fill, stroke=stroke, width=width))

    def polygon(self, points, fill=PALE, stroke=None, width=1):
        self.ops.append(dict(kind="polygon", points=points, fill=fill, stroke=stroke, width=width))

    def asset(self, name, x, y, w, h):
        path=ROOT/'results/figures/workflow/blueprint_design/details'/name
        self.ops.append(dict(kind="image",path=str(path),x=x,y=y,w=w,h=h))

    def card(self, x, y, w, h, title, body, color=TEAL, fill=PALE, number=None):
        self.rect(x, y, w, h, fill, color, .9)
        if number:
            self.text(x+14, y+10, w-28, number, 11, color, bold=True)
        self.text(x+14, y+(34 if number else 12), w-28, title, 20, color, bold=True)
        self.text(x+14, y+(72 if number else 48), w-28, body, 15, INK)

    def banner(self, y, text, fill=PALE, color=TEAL, size=17):
        self.rect(42, y, 876, 45, fill, None)
        self.text(56, y+9, 848, text, size, color)



def molecule(s,x,y,r=27,color=TEAL):
    pts=[(x+r*math.cos(k*math.pi/3),y+r*math.sin(k*math.pi/3)) for k in range(6)]
    s.line(pts+[pts[0]],color,2.4)
    for k in (0,2,4):
        a,b=pts[k],pts[(k+1)%6]
        a=(x+(a[0]-x)*.76,y+(a[1]-y)*.76);b=(x+(b[0]-x)*.76,y+(b[1]-y)*.76)
        s.line([a,b],color,1.5)
    for k in (0,2,4):
        p=pts[k];q=(x+(p[0]-x)*1.6,y+(p[1]-y)*1.6)
        s.line([p,q],color,2)
        s.ellipse(q[0]-5,q[1]-5,10,fill=GREEN if k==0 else GOLD,stroke=None)


def tree(s,x,y,w=62,h=68,color=TEAL):
    pts=[(x+w*.5,y),(x+w*.25,y+h*.45),(x+w*.75,y+h*.45),
         (x,y+h),(x+w*.35,y+h),(x+w*.65,y+h),(x+w,y+h)]
    for a,b in [(0,1),(0,2),(1,3),(1,4),(2,5),(2,6)]:s.line([pts[a],pts[b]],color,1.8)
    for j,(xx,yy) in enumerate(pts):s.ellipse(xx-4,yy-4,8,fill=WHITE if j<3 else color,stroke=color,width=1.3)


def lock(s,x,y,k=1,color=TEAL):
    s.ellipse(x+11*k,y,30*k,34*k,fill=None,stroke=color,width=2*k)
    s.rect(x,y+19*k,52*k,42*k,fill=PALE,stroke=color,width=1.5*k)
    s.ellipse(x+22*k,y+31*k,8*k,fill=color,stroke=None)
    s.line([(x+26*k,y+38*k),(x+26*k,y+48*k)],color,2*k)


def note_icon(s,x,y,w=57,h=71,color=TEAL):
    s.rect(x,y,w,h,WHITE,color,1.2)
    for j in range(4):s.line([(x+w*.18,y+h*(.23+j*.15)),(x+w*.81,y+h*(.23+j*.15))],color,1)


def infusion(s,x,y,k=1):
    s.line([(x+42*k,y-13*k),(x+42*k,y+93*k)],TEAL,1.5)
    s.rect(x,y,37*k,51*k,WHITE,TEAL,1.5)
    s.rect(x+4*k,y+28*k,29*k,18*k,PALE,None,rounded=False)
    s.line([(x+19*k,y+51*k),(x+19*k,y+69*k),(x+33*k,y+80*k),(x+33*k,y+100*k)],TEAL,1.5)
    s.line([(x+17*k,y-9*k),(x+17*k,y),(x+17*k,y-9*k),(x+42*k,y-9*k)],TEAL,1.2)
    s.ellipse(x+12*k,y+11*k,12*k,fill=GREEN_PALE,stroke=GREEN)


def curve(s,x,y,w,h,*,dots=False,initial=False,fill=False,labels=True,peak=.24):
    def cy(t):
        return .92*(1-math.exp(-6*t/peak))/(1-math.exp(-6)) if t<=peak else .92*math.exp(-3.6*(t-peak))
    pts=[(x+w*t,y+h*(1-cy(t))) for t in [k/70 for k in range(71)]]
    if fill:s.polygon([(x,y+h)]+pts+[(x+w,y+h)],fill=GREEN_PALE)
    s.line([(x,y-6),(x,y+h),(x+w+9,y+h)],LINE,1.2)
    if initial:
        old=[(x+w*t,y+h*(1-(.62*(1-math.exp(-5*t/peak))/(1-math.exp(-5)) if t<=peak else .62*math.exp(-1.8*(t-peak))))) for t in [k/50 for k in range(51)]]
        s.line(old,LINE,1.5,dash=True)
    s.line(pts,TEAL,2.5)
    if dots:
        for t,eps in [(.07,.02),(.2,-.045),(.38,.035),(.66,-.023),(.88,.019)]:
            px=x+w*t;py=y+h*(1-cy(t)+eps)
            s.ellipse(px-4,py-4,8,fill=WHITE,stroke=GOLD,width=2)
    if labels:
        s.eq(x-25,y-29,36,25,r'C_p',14,TEAL,role='axis label')
        s.text(x+w-8,y+h+4,25,'t',12,MUTED)
    return cy


def badge(s,x,y,text,color=TEAL,fill=PALE,w=150):
    s.rect(x,y,w,34,fill,None)
    s.text(x+8,y+5,w-16,text,15,color,align='center')


def explain(s, blocks, takeaway):
    """Readable on-slide prose, deliberately separate from technical speaker notes."""
    s.line([(585,145),(585,437)],'DCE5E8',1)
    for y,(title,body) in zip((149,249,349),blocks):
        s.text(609,y,309,title,fit_size(title,309,27,19,minimum=14),TEAL,bold=True,line=1.2)
        h=s.text(609,y+31,309,body,fit_size(body,309,69,17,line=1.28,minimum=12),INK,line=1.28)
        if h > 69:raise ValueError('Explanation too long: '+body)
    s.rect(42,457,876,32,PALE,None)
    s.text(54,462,852,takeaway,14,TEAL,line=1.25)


def make_slides():
    from blueprint_chain_slides import make_blueprint_chain_slides
    return make_blueprint_chain_slides()


def equation_assets(op):
    fig=Figure(figsize=(14,2),dpi=250)
    FigureCanvasAgg(fig)
    fig.patch.set_alpha(0)
    obj=fig.text(.015,.44,"$"+op["latex"]+"$",fontsize=op["size"],
                 color="#"+op["color"],va="center",math_fontfamily="stix")
    fig.canvas.draw()
    bbox=obj.get_window_extent(fig.canvas.get_renderer()).expanded(1.018,1.09)
    crop=bbox.transformed(fig.dpi_scale_trans.inverted())
    data=[]
    for fmt in ("png","pdf"):
        bio=io.BytesIO()
        fig.savefig(bio,format=fmt,dpi=450,transparent=True,bbox_inches=crop,pad_inches=0)
        data.append(bio.getvalue())
    return data


def add_arrow_xml(line, end="tailEnd"):
    el=OxmlElement("a:"+end)
    el.set("type","triangle");el.set("w","med");el.set("len","med")
    line.get_or_add_ln().append(el)


def pdf_arrow(page,a,b,color,width):
    dx,dy=b[0]-a[0],b[1]-a[1]
    length=math.hypot(dx,dy)
    if length==0:return
    ux,uy=dx/length,dy/length
    size=max(5,3.8*width)
    pts=[b,(b[0]-ux*size-uy*size*.48,b[1]-uy*size+ux*size*.48),
         (b[0]-ux*size+uy*size*.48,b[1]-uy*size-ux*size*.48),b]
    page.draw_polyline(pts,color=rgb(color),fill=rgb(color),width=.1,closePath=True)


def render(slides, *, output_dir=None, stem=None, source_pptx=None, source_pdf=None,
           title='PBPK blueprint design: principles, inputs, computation, and outputs',
           subject='Physiological equations, evidence-constrained calibration, feasibility checks, and locked prediction',
           required_terms=None, source_indices=None, prepend=True, toc=None):
    output_dir=Path(output_dir or OUT); stem=stem or STEM
    output_dir.mkdir(parents=True,exist_ok=True)
    if bool(source_pptx) != bool(source_pdf):
        raise ValueError('A preserved deck requires both PPTX and PDF sources')
    prs=Presentation(source_pptx) if source_pptx else Presentation()
    source_count=len(prs.slides)
    if source_indices is not None:
        keep=set(source_indices)
        for i,element in reversed(list(enumerate(list(prs.slides._sldIdLst)))):
            if i not in keep:
                prs.part.drop_rel(element.rId)
                prs.slides._sldIdLst.remove(element)
    if source_pptx and (prs.slide_width != Pt(W) or prs.slide_height != Pt(H)):
        raise ValueError('Source slide dimensions differ')
    prs.slide_width=Pt(W);prs.slide_height=Pt(H)
    original_count=len(prs.slides)
    original_titles=[]
    for slide in prs.slides:
        original_titles.append(next((sh.text.replace('\n',' ') for sh in slide.shapes
            if sh.has_text_frame and abs(sh.top/Pt(1)-55)<1), ''))
    total_count=original_count+len(slides)
    prs.core_properties.title=title
    prs.core_properties.subject=subject
    prs.core_properties.author="PBPK research project"
    doc=fitz.open()
    def insert_preserved():
        with fitz.open(source_pdf) as preserved:
            if len(preserved)!=source_count:
                raise ValueError('PPTX and PDF source page counts differ')
            if source_indices is not None:preserved.select(source_indices)
            doc.insert_pdf(preserved)
    if original_count and not prepend:insert_preserved()
    counts={"slides":total_count,"preserved_slides":original_count,"editable_text_boxes":0,"equations":0,"shapes":0}
    for item in slides:
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)
        # Speaker notes are intentionally absent from the delivered deck.
        page=doc.new_page(width=W,height=H)
        # Keep the CID-keyed CJK fonts intact so glyph IDs remain stable in PDF.
        page.insert_font(fontname="cjk",fontfile=str(FONT_FILE))
        page.insert_font(fontname="cjkb",fontfile=str(BOLD_FILE))
        for op in item.ops:
            kind=op["kind"]
            if kind=="text":
                shape=slide.shapes.add_textbox(Pt(op['x']),Pt(op['y']),Pt(op['w']),Pt(op['h']))
                tf=shape.text_frame;tf.clear();tf.word_wrap=False
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                tf.vertical_anchor=MSO_ANCHOR.TOP
                for j,txt in enumerate(op["rows"]):
                    p=tf.paragraphs[0] if j==0 else tf.add_paragraph()
                    p.alignment={"left":PP_ALIGN.LEFT,"center":PP_ALIGN.CENTER,"right":PP_ALIGN.RIGHT}[op["align"]]
                    p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(op["size"]*op["spacing"])
                    run=p.add_run();run.text=txt;f=run.font
                    face=BOLD_NAME if op['bold'] else FONT_NAME
                    f.name=face;f.size=Pt(op['size']);f.bold=op['bold'];f.color.rgb=RGBColor.from_string(op['color'])
                    rpr=run._r.get_or_add_rPr()
                    for tag in ("a:ea","a:cs"):
                        node=OxmlElement(tag);node.set("typeface",face);rpr.append(node)
                    width=(BOLD_FONT if op['bold'] else FONT).text_length(txt,fontsize=op["size"])
                    xx=op['x']+({"left":0,"center":(op['w']-width)/2,"right":op['w']-width}[op['align']])
                    yy=op['y']+op['size']*1.12+j*op['size']*op['spacing']
                    page.insert_text((xx,yy),txt,fontname="cjkb" if op['bold'] else "cjk",
                                     fontsize=op['size'],color=rgb(op['color']))
                counts['editable_text_boxes']+=1
            elif kind=="rect":
                shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if op['rounded'] else MSO_SHAPE.RECTANGLE,
                                             Pt(op['x']),Pt(op['y']),Pt(op['w']),Pt(op['h']))
                if op['rounded']:shape.adjustments[0]=.05
                shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(op['fill'])
                if op['stroke']:
                    shape.line.color.rgb=RGBColor.from_string(op['stroke']);shape.line.width=Pt(op['width'])
                else:shape.line.fill.background()
                page.draw_rect(fitz.Rect(op['x'],op['y'],op['x']+op['w'],op['y']+op['h']),
                               color=rgb(op['stroke']) if op['stroke'] else None,fill=rgb(op['fill']),
                               width=op['width'],radius=.05 if op['rounded'] else None)
                counts['shapes']+=1
            elif kind in {"ellipse", "polygon", "image"}:
                if kind=="ellipse":
                    shape=slide.shapes.add_shape(MSO_SHAPE.OVAL,Pt(op['x']),Pt(op['y']),Pt(op['w']),Pt(op['h']))
                    if op['fill']:
                        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(op['fill'])
                    else:shape.fill.background()
                    if op['stroke']:
                        shape.line.color.rgb=RGBColor.from_string(op['stroke']);shape.line.width=Pt(op['width'])
                    else:shape.line.fill.background()
                    page.draw_oval(fitz.Rect(op['x'],op['y'],op['x']+op['w'],op['y']+op['h']),
                                   color=rgb(op['stroke']) if op['stroke'] else None,
                                   fill=rgb(op['fill']) if op['fill'] else None,width=op['width'])
                elif kind=="polygon":
                    pts=op['points']
                    builder=slide.shapes.build_freeform(Pt(pts[0][0]),Pt(pts[0][1]))
                    builder.add_line_segments([(Pt(x),Pt(y)) for x,y in pts[1:]],close=True)
                    shape=builder.convert_to_shape()
                    if op['fill']:
                        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(op['fill'])
                    else:shape.fill.background()
                    if op['stroke']:
                        shape.line.color.rgb=RGBColor.from_string(op['stroke']);shape.line.width=Pt(op['width'])
                    else:shape.line.fill.background()
                    page.draw_polyline(pts,color=rgb(op['stroke']) if op['stroke'] else None,
                                       fill=rgb(op['fill']) if op['fill'] else None,
                                       closePath=True,width=op['width'])
                else:
                    im=Image.open(op['path']);iw,ih=im.size
                    scale=min(op['w']/iw,op['h']/ih);w,h=iw*scale,ih*scale
                    x,y=op['x']+(op['w']-w)/2,op['y']+(op['h']-h)/2
                    shape=slide.shapes.add_picture(op['path'],Pt(x),Pt(y),Pt(w),Pt(h))
                    if op.get('pdf'):
                        with fitz.open(op['pdf']) as vector_image:
                            page.show_pdf_page(fitz.Rect(x,y,x+w,y+h),vector_image,0)
                    else:
                        page.insert_image(fitz.Rect(x,y,x+w,y+h),filename=op['path'])
                counts['shapes']+=1
            elif kind=="line":
                points=op['points']
                for j,(a,b) in enumerate(zip(points,points[1:])):
                    shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(a[0]),Pt(a[1]),Pt(b[0]),Pt(b[1]))
                    shape.line.color.rgb=RGBColor.from_string(op['color']);shape.line.width=Pt(op['width'])
                    if op['dash']:
                        dash=OxmlElement("a:prstDash");dash.set("val","dash");shape.line._get_or_add_ln().append(dash)
                    if op['arrow'] and j==len(points)-2:add_arrow_xml(shape._element.spPr,"tailEnd")
                    if op['both'] and j==0:add_arrow_xml(shape._element.spPr,"headEnd")
                page.draw_polyline(points,color=rgb(op['color']),width=op['width'],dashes="[5 4] 0" if op['dash'] else None)
                if op['arrow']:pdf_arrow(page,points[-2],points[-1],op['color'],op['width'])
                if op['both']:pdf_arrow(page,points[1],points[0],op['color'],op['width'])
            else:
                if op.get('role', 'formula') == 'formula':
                    assert op['size'] == EQUATION_FONT_SIZE, 'Use the shared equation font size'
                png,pdf=equation_assets(op)
                with fitz.open(stream=pdf,filetype="pdf") as eqdoc:
                    # PDF coordinates are points: retain the natural typeset size.
                    # Fitting each image to its box made short formulas much larger.
                    w,h=eqdoc[0].rect.width,eqdoc[0].rect.height
                    if w>op['w']+.1 or h>op['h']+.1:
                        raise ValueError(f"Reflow formula instead of resizing it: {op['latex']}")
                    x,y=op['x']+(op['w']-w)/2,op['y']+(op['h']-h)/2
                    shape=slide.shapes.add_picture(io.BytesIO(png),Pt(x),Pt(y),Pt(w),Pt(h))
                    shape.name=f"{op.get('role', 'formula')} {op['size']:g} pt"
                    shape._element.nvPicPr.cNvPr.set("descr",op['latex'])
                    page.show_pdf_page(fitz.Rect(x,y,x+w,y+h),eqdoc,0)
                counts['equations']+=1
    if original_count and prepend:
        # Append natively, then reorder slide relationships; existing shapes,
        # pictures, equations and notes survive without cloning or flattening.
        ids=list(prs.slides._sldIdLst)
        for element in ids:
            prs.slides._sldIdLst.remove(element)
        for element in ids[original_count:]+ids[:original_count]:
            prs.slides._sldIdLst.append(element)
        insert_preserved()
    if original_count:
        preserved_range=range(len(slides),total_count) if prepend else range(original_count)
        for index in preserved_range:
            label=f'{index+1:02d} / {total_count:02d}'
            matches=[sh for sh in prs.slides[index].shapes if sh.has_text_frame
                     and re.fullmatch(r'\d+ / \d+',sh.text.strip())]
            if len(matches)!=1:raise ValueError('Missing or ambiguous original footer')
            runs=matches[0].text_frame.paragraphs[0].runs
            if len(runs)!=1:raise ValueError('Unexpected footer run structure')
            runs[0].text=label
            page=doc[index]
            page.add_redact_annot(fitz.Rect(850,500,921,523),fill=(1,1,1))
            page.apply_redactions(images=0,graphics=0,text=0)
            page.insert_font(fontname='newfooter',fontfile=str(FONT_FILE))
            x=918-FONT.text_length(label,fontsize=10)
            page.insert_text((x,514.2),label,fontname='newfooter',fontsize=10,color=rgb(TEAL))
    # Drop the relationships as well as text, so old notes XML is not shipped.
    for slide in prs.slides:
        for rel in list(slide.part.rels.values()):
            if rel.reltype.endswith('/notesSlide'):slide.part.drop_rel(rel.rId)
    for rel in list(prs.part.rels.values()):
        if rel.reltype.endswith('/notesMaster'):prs.part.drop_rel(rel.rId)
    for element in list(prs._element):
        if element.tag.endswith('}notesMasterIdLst'):prs._element.remove(element)
    pptx_path=output_dir/(stem+".pptx");pdf_path=output_dir/(stem+".pdf")
    prs.save(pptx_path)
    doc.set_metadata({"title":prs.core_properties.title,"subject":prs.core_properties.subject,
                      "author":"PBPK research project","creator":"Project-local editable PPTX / PDF builder"})
    if toc:
        doc.set_toc(toc)
    doc.save(pdf_path,garbage=4,deflate=True)
    thumbs=[]
    for i,page in enumerate(doc):
        pix=page.get_pixmap(matrix=fitz.Matrix(1.0,1.0),alpha=False)
        img=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
        img.thumbnail((720,405))
        thumbs.append(img)
    contact=Image.new('RGB',(1480,10+430*math.ceil(total_count/2)),'#E7ECEF')
    draw=ImageDraw.Draw(contact)
    font=ImageFont.truetype(str(FONT_FILE),15)
    new_titles=[item.ops[1]['rows'][0] for item in slides]
    headings=new_titles+original_titles if prepend else original_titles+new_titles
    for i,img in enumerate(thumbs):
        x=10+(i%2)*740;y=10+(i//2)*430
        heading=headings[i]
        contact.paste(img,(x,y));draw.text((x+8,y+407),f"{i+1:02d}  {heading}",font=font,fill='#255E72')
    preview=output_dir/(stem+'_preview.png');contact.save(preview)
    doc.close()
    # Artifact integrity, actual XML content and readable-page checks.
    check=Presentation(pptx_path)
    assert len(check.slides)==total_count
    with zipfile.ZipFile(pptx_path) as z:
        assert z.testzip() is None
        assert not any(n.startswith(('ppt/notesSlides/','ppt/notesMasters/')) for n in z.namelist())
    with fitz.open(pdf_path) as final:
        assert len(final)==total_count
        text=unicodedata.normalize('NFKC','\n'.join(p.get_text() for p in final))
        text=text.replace('\u2010','-').replace('\u2011','-')
        for term in required_terms or ['Human','Differential equations','Calibration','Evidence']:
            assert term in text,term
        for page in final:
            for block in page.get_text('dict')['blocks']:
                if block.get('type')==0:
                    for line in block['lines']:
                        for span in line['spans']:
                            x0,y0,x1,y1=span['bbox']
                            assert x0>=-1 and y0>=-1 and x1<=W+1 and y1<=H+1,span
    print(counts)
    for path in (pptx_path,pdf_path,preview):print(path.relative_to(ROOT),path.stat().st_size)


def module_boundary_slide():
    s=Slide('Selection networks choose models; blueprint equations predict PK',
        'The earlier neural networks support evidence and model selection. The current blueprint does not predict concentrations with a neural network.',
        'MODULE BOUNDARY','')
    for j,(x,n,col) in enumerate([(88,5,TEAL),(238,4,GREEN),(394,2,GOLD)]):
        for k in range(n):
            yy=163+k*30+(5-n)*15
            s.ellipse(x,yy,13,fill=WHITE,stroke=col)
            if j<2:
                nx,nn,ncol=[(238,4,GREEN),(394,2,GOLD)][j]
                for kk in range(nn):
                    ny=163+kk*30+(5-nn)*15
                    s.line([(x+14,yy+6),(nx-2,ny+6)],LINE,.55)
    s.text(47,320,167,'716-dimensional input',15,TEAL,align='center')
    s.text(215,320,145,'128-dimensional hidden layer',15,GREEN,align='center')
    s.text(376,320,174,'Ranking and other heads',15,GOLD,align='center')
    s.eq(52,346,512,34,r'\mathbf{x}=[\mathbf{u}_Q;\mathbf{u}_m;|\mathbf{u}_Q-\mathbf{u}_m|;\mathbf{u}_Q\odot\mathbf{u}_m;\mathbf{f}]',EQUATION_FONT_SIZE,TEAL)
    s.eq(52,384,512,33,r'\mathbf{h}=\mathrm{GELU}(\mathrm{LN}(W\mathbf{x}+\mathbf{b}))',EQUATION_FONT_SIZE,TEAL)
    s.eq(52,422,512,33,r's_{\rm doc}=\mathbf{w}_s^{\mathsf{T}}\mathbf{h}+b_s',EQUATION_FONT_SIZE,TEAL)
    explain(s,[('Network inputs: requests and candidates','Each text is encoded into 128 dimensions. Differences, products, and 204 condition features form a 716-dimensional input.'),
               ('Network outputs: selection scores','An MLP uses linear transformations, normalization, and activation functions to calculate scores, not PK curves.'),
               ('A blueprint follows when reuse is unavailable','Hypothetical drug X has evidence and a few development observations but no reusable model. Equations then support construction and calibration.')],
            'Handoff: no suitable candidate + drug evidence and request → physiology, equation calibration, and target prediction.')
    return s


def refresh_current_presentation():
    """Publish the introductory teaching sequence and optional technical references."""
    import io
    import tempfile
    from blueprint_chain_slides import make_blueprint_technical_slides
    from presentation_teaching_common import make_foundation_slides, make_teaching_conclusion_slides
    pptx_path,pdf_path=OUT/(STEM+'.pptx'),OUT/(STEM+'.pdf')
    old_ppt=pptx_path.read_bytes();old_pdf=pdf_path.read_bytes()
    Presentation(io.BytesIO(old_ppt))  # Verify the source before replacing the requested deck.
    groups=[('PK and PBPK from first principles',make_foundation_slides()[:4]),
            ('Blueprint design: construction, computation, and validation',make_slides()),
            ('Blueprint results and workflow recap',make_teaching_conclusion_slides())]
    main_count=sum(len(pages) for _,pages in groups)
    groups += [('Technical reference: PBPK equations and endpoints',make_blueprint_technical_slides())]
    fresh=[];toc=[]
    for heading,pages in groups:
        toc.append([1,heading,len(fresh)+1])
        fresh.extend(pages)
    total=len(fresh)
    assert main_count==22 and total==27
    for i,page in enumerate(fresh,1):
        assert len(page.ops[1]['rows'])==1, 'Shorten title: '+str(page.ops[1]['rows'])
        assert len(page.ops[2]['rows'])==1, 'Shorten subtitle: '+str(page.ops[2]['rows'])
        if i>main_count:
            page.ops[4]['rows']=[f'Technical reference · Optional on first reading · Main slides 1–{main_count}']
        page.text(855,503,63,f'{i:02d} / {total:02d}',10,TEAL,align='right')
        page.notes=''
    work_root=ROOT/'runtime/presentation_build';work_root.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(dir=work_root,prefix='beginner_teaching-') as tmp:
        work=Path(tmp)
        render(fresh,output_dir=work,
               title='Understanding PBPK blueprints: principles and workflow',toc=toc,
               required_terms=['Blueprint design','Differential equations','Calibration','Teaching','Human'])
        new_ppt,new_pdf=work/(STEM+'.pptx'),work/(STEM+'.pdf')
        updated=Presentation(new_ppt)
        assert len(updated.slides)==total
        for i,page in enumerate(updated.slides):
            assert not page.has_notes_slide,f'Notes remain: {i+1}'
            for sh in page.shapes:
                assert sh.left>=-12700 and sh.top>=-12700
                assert sh.left+sh.width<=updated.slide_width+12700
                assert sh.top+sh.height<=updated.slide_height+12700
        blueprint_text='\n'.join(sh.text for page in updated.slides
                                for sh in page.shapes if sh.has_text_frame)
        assert not re.search('(?i)component recommendation|model recommendation|random forest|\u7ec4\u4ef6\u63a8\u8350|\u6a21\u578b\u63a8\u8350|\\bRF\\b|\u968f\u673a\u68ee\u6797|main\\.tex|permeability_model\\.py',blueprint_text)
        formula_sizes={op['size'] for page in fresh for op in page.ops
                       if op['kind']=='equation' and op.get('role','formula')=='formula'}
        assert formula_sizes=={EQUATION_FONT_SIZE}
        if pptx_path.read_bytes()!=old_ppt or pdf_path.read_bytes()!=old_pdf:
            raise RuntimeError('Presentation changed during build; refusing to overwrite concurrent edits')
        for suffix in ['.pptx','.pdf','_preview.png']:
            (work/(STEM+suffix)).replace(OUT/(STEM+suffix))
    print({'total_slides':total,'beginner_main_slides':main_count,'technical_reference_slides':total-main_count,
           'formula_font_size_pt':EQUATION_FONT_SIZE,'notes_parts':0,'pdf_navigation_sections':len(toc)})


if __name__=="__main__":
    refresh_current_presentation()
