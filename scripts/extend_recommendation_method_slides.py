#!/usr/bin/env python3
"""Prepend two editable method sections while retaining all blueprint slides."""
from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
import io
import json
from pathlib import Path
import re
import tempfile
import zipfile

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
import pymupdf as fitz

import build_blueprint_method_slides as style
from component_method_slides import make_component_slides
from model_method_slides import make_model_slides

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'results/presentations'
STEM=style.STEM
WORK=ROOT/'runtime/presentation_build/recommendation_methods'


def texts(slide):
    return [shape.text for shape in slide.shapes if shape.has_text_frame]


def original_indices(prs):
    result=[i for i,slide in enumerate(prs.slides)
        if any(t.startswith('PBPK BLUEPRINT   /') for t in texts(slide))]
    if len(result)!=10:
        raise ValueError('Expected all ten existing blueprint slides; refusing to discard content')
    for i,slide in enumerate(prs.slides):
        if i not in result and not any(t.startswith(('PBPK COMPONENTS   /','PBPK MODEL RECOMMENDATION   /')) for t in texts(slide)):
            raise ValueError('Unrecognized user slide; preserve it before rebuilding')
    return result


def without_footer_xml(slide):
    el=deepcopy(slide._element)
    for t in el.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t'):
        if re.fullmatch(r'\d+ / \d+',t.text or ''):t.text='PAGE / TOTAL'
    return etree.tostring(el)


def save_base(pptx_data,pdf_data,work):
    prs=Presentation(io.BytesIO(pptx_data))
    keep=original_indices(prs)
    ids=list(prs.slides._sldIdLst)
    for i,el in reversed(list(enumerate(ids))):
        if i not in keep:
            prs.part.drop_rel(el.rId)
            prs.slides._sldIdLst.remove(el)
    source=work/'preserved_blueprint.pptx';prs.save(source)
    with fitz.open(stream=pdf_data,filetype='pdf') as doc:
        if len(doc)!=len(ids):raise ValueError('Original PPTX and PDF are out of sync')
        doc.select(keep)
        doc.save(work/'preserved_blueprint.pdf')
    return source,work/'preserved_blueprint.pdf'


def validate(old_ppt,old_pdf,new_ppt,new_pdf):
    before,after=Presentation(old_ppt),Presentation(new_ppt)
    assert len(before.slides)==10 and len(after.slides)==30
    for i,old in enumerate(before.slides):
        new=after.slides[20+i]
        assert without_footer_xml(old)==without_footer_xml(new),f'Blueprint slide {i+1} changed'
        assert old.notes_slide.notes_text_frame.text==new.notes_slide.notes_text_frame.text
        old_images=[sha256(sh.image.blob).hexdigest() for sh in old.shapes if sh.shape_type==13]
        new_images=[sha256(sh.image.blob).hexdigest() for sh in new.shapes if sh.shape_type==13]
        assert old_images==new_images
    assert all('PBPK COMPONENTS' in texts(after.slides[i])[0] for i in range(10))
    assert all('PBPK MODEL RECOMMENDATION' in texts(after.slides[i])[0] for i in range(10,20))
    for i,slide in enumerate(after.slides):
        assert f'{i+1:02d} / 30' in texts(slide)
        for sh in slide.shapes:
            assert sh.left>=-12700 and sh.top>=-12700
            assert sh.left+sh.width<=after.slide_width+12700
            assert sh.top+sh.height<=after.slide_height+12700
    with fitz.open(old_pdf) as old,fitz.open(new_pdf) as new:
        assert len(new)==30
        for i in range(10):
            clip=fitz.Rect(0,0,960,495)
            a=old[i].get_pixmap(matrix=fitz.Matrix(1,1),clip=clip,alpha=False)
            b=new[20+i].get_pixmap(matrix=fitz.Matrix(1,1),clip=clip,alpha=False)
            assert a.samples==b.samples,f'Blueprint visual content changed: {i+1}'
    with zipfile.ZipFile(new_ppt) as archive:
        assert archive.testzip() is None


def previews(pdf,work):
    font=ImageFont.truetype(str(style.FONT_FILE),15)
    title_font=ImageFont.truetype(str(style.FONT_FILE),19)
    with fitz.open(pdf) as doc:
        all_images=[]
        for page in doc:
            pix=page.get_pixmap(matrix=fitz.Matrix(1,1),alpha=False)
            all_images.append(Image.frombytes('RGB',(pix.width,pix.height),pix.samples))
        overview=Image.new('RGB',(1490,50+10*291),'#E7ECEF')
        draw=ImageDraw.Draw(overview)
        for section,title in enumerate(['Component recommendation · 01—10','Model recommendation · 11—20','Blueprint design · 21—30']):
            draw.text((20+section*490,12),title,font=title_font,fill='#255E72')
            contact=Image.new('RGB',(1480,2160),'#E7ECEF')
            cd=ImageDraw.Draw(contact)
            for local in range(10):
                index=section*10+local
                img=all_images[index].copy();img.thumbnail((480,270))
                overview.paste(img,(10+section*490,50+local*291))
                draw.text((18+section*490,320+local*291),f'{index+1:02d}',font=font,fill='#255E72')
                larger=all_images[index].copy();larger.thumbnail((720,405))
                x=10+(local%2)*740;y=10+(local//2)*430
                contact.paste(larger,(x,y));cd.text((x+8,y+407),f'{index+1:02d}',font=font,fill='#255E72')
            contact.save(work/f'preview_section_{section+1}.png')
        overview.save(work/(STEM+'_preview.png'))


def main():
    WORK.mkdir(parents=True,exist_ok=True)
    pptx_path,pdf_path=OUT/(STEM+'.pptx'),OUT/(STEM+'.pdf')
    pptx_data,pdf_data=pptx_path.read_bytes(),pdf_path.read_bytes()
    base_ppt,base_pdf=save_base(pptx_data,pdf_data,WORK)
    slides=make_component_slides()+make_model_slides()
    assert len(slides)==20
    for i,slide in enumerate(slides):
        slide.text(855,503,63,f'{i+1:02d} / 30',10,style.TEAL,align='right')
        slide.notes+=' Graphics explain the methods. Numerical results come from artifacts listed in the notes; other nodes, weights, and workflows are illustrative.'
    style.render(slides,output_dir=WORK,stem=STEM,source_pptx=base_ppt,source_pdf=base_pdf,
        title='PBPK component recommendation, model recommendation, and blueprint design: principles and workflow',
        subject='Request–evidence graph matching, PBPK-Match, and blueprint construction with fixed physiology',
        required_terms=['Component recommendation','Qwen3-Embedding','204','716','280','0.909','Human','RF','Calibration'])
    new_ppt,new_pdf=WORK/(STEM+'.pptx'),WORK/(STEM+'.pdf')
    validate(base_ppt,base_pdf,new_ppt,new_pdf)
    previews(new_pdf,WORK)
    result={'slides':30,'added_component_slides':10,'added_model_slides':10,'preserved_blueprint_slides':10,
        'blueprint_slide_xml_unchanged_except_page_numbers':True,'blueprint_notes_and_images_preserved':True,
        'blueprint_pdf_pixels_unchanged_above_footer':True,
        'original_pptx_sha256':sha256(pptx_data).hexdigest(),'original_pdf_sha256':sha256(pdf_data).hexdigest(),
        'final_pptx_sha256':sha256(new_ppt.read_bytes()).hexdigest(),
        'final_pdf_sha256':sha256(new_pdf.read_bytes()).hexdigest()}
    (WORK/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    # Publish only after structural, source-preservation and rendered-page checks.
    if pptx_path.read_bytes()!=pptx_data or pdf_path.read_bytes()!=pdf_data:
        raise RuntimeError('Presentation changed during build; do not overwrite concurrent edits')
    for suffix in ['.pptx','.pdf','_preview.png']:
        src=WORK/(STEM+suffix);dst=OUT/(STEM+suffix)
        with tempfile.NamedTemporaryFile(dir=OUT,prefix='.presentation-',delete=False) as f:
            temp=Path(f.name);f.write(src.read_bytes())
        temp.replace(dst)
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':
    main()
