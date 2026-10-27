"""Three research-framework designs: native PPTX objects, SVG and PNG previews.

No screenshots or generated bitmaps are embedded in the deliverable slides.
The imagegen concept is a visual reference only. All diagrams use this shared
scene representation, so labels, arrows and biological icons remain editable.
Run with the project runtime/app-env/bin/python; requires Pillow and python-pptx.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from math import atan2, cos, sin, pi
import json
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'results/figures/workflow'
FONT = ROOT / 'runtime/fonts/NotoSansCJKsc-Regular.otf'
W, H, SCALE = 1920, 1080, 2
INK, MUTED, EDGE = '#233746', '#65727C', '#87949C'
WHITE = '#FFFFFF'
PALETTES = [
    ('#EDF4FA', '#ABC3D6', '#426C8C'),
    ('#EDF5EE', '#B8CEBB', '#52795D'),
    ('#FCF6E5', '#DFCEA0', '#9B7B34'),
    ('#F2EEF8', '#C7BCD9', '#7D679A'),
]
SVG_NS = 'http://www.w3.org/2000/svg'
ET.register_namespace('', SVG_NS)


@dataclass
class Primitive:
    kind: str
    args: dict


class Figure:
    def __init__(self, slug, title):
        self.slug, self.title = slug, title
        self.items: list[Primitive] = []
        self.text_bounds = []
        self.label_context = ''
        self.rect(0, 0, W, H, '#FFFFFF', None, 0)

    def add(self, kind, **kwargs):
        kwargs['name'] = self.label_context + '/' + kind + f'/{len(self.items):03d}'
        self.items.append(Primitive(kind, kwargs))

    def rect(self, x, y, w, h, fill=WHITE, stroke=EDGE, radius=12, width=1.6):
        self.add('rect', x=x, y=y, w=w, h=h, fill=fill, stroke=stroke, radius=radius, width=width)

    def ellipse(self, x, y, w, h, fill=WHITE, stroke=EDGE, width=1.8):
        self.add('ellipse', x=x, y=y, w=w, h=h, fill=fill, stroke=stroke, width=width)

    def polygon(self, pts, fill, stroke=None, width=1.6, arrowhead=False):
        self.add('polygon', pts=pts, fill=fill, stroke=stroke, width=width, arrowhead=arrowhead)

    def line(self, pts, color=EDGE, width=2, arrow=False, dashed=False):
        self.add('line', pts=pts, color=color, width=width, dashed=dashed, arrow=arrow)
        if arrow:
            x, y = pts[-1]
            a = atan2(y-pts[-2][1], x-pts[-2][0])
            self.polygon([(x,y), (x-10*cos(a)+4*sin(a), y-10*sin(a)-4*cos(a)),
                          (x-10*cos(a)-4*sin(a), y-10*sin(a)+4*cos(a))], color, arrowhead=True)

    def text(self, x, y, w, h, value, size=24, color=INK, align='left', bold=False):
        font = ImageFont.truetype(str(FONT), round(size*SCALE))
        for line in value.split('\n'):
            tw = font.getlength(line)/SCALE
            if tw > w-3:
                raise ValueError(f'{self.slug}: text width {tw:.1f}>{w}: {line}')
        if len(value.split('\n'))*size*1.28 > h+1:
            raise ValueError(f'{self.slug}: text height: {value!r} ({h})')
        if min(x,y)<0 or x+w>W+1 or y+h>H+1:
            raise ValueError(f'{self.slug}: text outside canvas: {value}')
        self.text_bounds.append([x,y,w,h,value])
        self.add('text', x=x, y=y, w=w, h=h, value=value, size=size, color=color, align=align, bold=bold)

    def heading(self, variant):
        self.text(48, 24, 1810, 66, '证据驱动的 PBPK 建模推荐框架', 46, bold=True)
        self.text(51, 96, 1300, 35, '资料与知识库  →  组件推荐  →  模型推荐  →  蓝图设计', 23, MUTED)
        self.text(1570, 101, 296, 30, variant, 20, MUTED, 'right')
        self.line([(48,148),(1872,148)], '#DDE3E7', 1.4)

    def footer(self):
        self.line([(48,1024),(1872,1024)], '#DDE3E7', 1.3)
        self.text(51,1034,1810,35,
                  '全程证据溯源：原始资料 → 证据引用 → 推荐依据 → 参数来源    |    蓝图输出 → 仿真、校准与独立验证',
                  21,MUTED)

    def panel(self, x,y,w,h,letter,title,index,subtitle=None):
        pale, border, accent = PALETTES[index]
        self.label_context = letter+' '+title
        self.rect(x,y,w,h,pale,border,19,1.8)
        self.text(x+23,y+17,50,43,f'({letter})',29,accent,bold=True)
        self.text(x+85,y+17,w-110,43,title,29,INK,bold=True)
        if subtitle:
            self.text(x+26,y+67,w-52,31,subtitle,20,MUTED)
        else:
            self.line([(x+23,y+69),(x+w-23,y+69)],border,1)


def icon(f, kind, x, y, size, color, pale=None):
    """Semantic line illustrations built exclusively from editable primitives."""
    pale = pale or '#F7FAFC'
    def p(a,b): return (x+a*size,y+b*size)
    def ln(points, c=color, width=1.8, arrow=False, dashed=False):
        f.line([p(a,b) for a,b in points], c, width, arrow, dashed)
    def rr(a,b,w,h,fill=WHITE,stroke=color,r=4):
        f.rect(x+a*size,y+b*size,w*size,h*size,fill,stroke,r)
    def el(a,b,w,h,fill=WHITE,stroke=color):
        f.ellipse(x+a*size,y+b*size,w*size,h*size,fill,stroke)
    def poly(pts,fill=pale,stroke=color):
        f.polygon([p(a,b) for a,b in pts],fill,stroke)

    if kind == 'documents':
        rr(.12,.18,.57,.68,pale)
        rr(.22,.1,.57,.7,WHITE)
        poly([(.34,.03),(.75,.03),(.92,.2),(.92,.73),(.34,.73)],WHITE)
        ln([(.75,.03),(.75,.2),(.92,.2)])
        for yy in (.3,.42,.54): ln([(.43,yy),(.8,yy)],width=1.5)
        el(.03,.64,.24,.24,pale)
        ln([(.23,.84),(.36,.97)],width=3)
    elif kind == 'extract':
        rr(.05,.14,.37,.67)
        for yy in (.29,.42,.55,.68): ln([(.11,yy),(.34,yy)],'#ABB6BD',1.5)
        ln([(.44,.45),(.62,.45)],arrow=True)
        for yy in (.15,.4,.65):
            rr(.7,yy,.24,.17,pale)
        ln([(.66,.23),(.63,.23),(.63,.73),(.66,.73)],width=1.3)
    elif kind == 'quality':
        poly([(.5,.06),(.88,.19),(.84,.58),(.7,.79),(.5,.93),(.3,.79),(.16,.58),(.12,.19)],pale)
        ln([(.28,.45),(.45,.63),(.74,.32)],width=3)
    elif kind in ('graph','retrieval'):
        nodes=[(.5,.45),(.18,.13),(.82,.14),(.15,.76),(.81,.79),(.52,.96)]
        for a,b in ((0,1),(0,2),(0,3),(0,4),(0,5),(1,3),(2,4),(3,5),(4,5)):
            ln([nodes[a],nodes[b]],'#A5B5BD',1.5)
        for i,(a,b) in enumerate(nodes):
            el(a-.085,b-.085,.17,.17,color if i==0 else pale,color)
        if kind=='retrieval':
            el(.47,.43,.4,.4,None,color)
            ln([(.82,.78),(.99,.98)],width=4)
    elif kind == 'query':
        rr(.05,.07,.88,.67,WHITE,r=8)
        poly([(.22,.73),(.22,.94),(.44,.73)],WHITE)
        el(.17,.21,.2,.2,pale)
        ln([(.13,.57),(.16,.48),(.39,.48),(.43,.57)])
        for yy in (.26,.41,.57): ln([(.54,yy),(.82,yy)],width=2)
    elif kind == 'components':
        for i,(a,b) in enumerate(((.08,.12),(.62,.12),(.08,.65),(.62,.65))):
            rr(a,b,.3,.28,PALETTES[i][0],color,r=6)
        ln([(.38,.26),(.61,.26)],arrow=True)
        ln([(.23,.41),(.23,.64)],arrow=True)
        ln([(.76,.41),(.76,.64)],arrow=True)
        ln([(.39,.79),(.61,.79)],arrow=True)
    elif kind == 'model':
        rr(.06,.04,.88,.89,WHITE,r=7)
        for a,b in ((.21,.23),(.61,.23),(.42,.56)):
            rr(a,b,.19,.16,pale,r=2)
        ln([(.4,.31),(.6,.31)],width=1.7,arrow=True)
        ln([(.3,.4),(.3,.64),(.42,.64)],width=1.7,arrow=True)
        ln([(.7,.4),(.7,.64),(.63,.64)],width=1.7,arrow=True)
        ln([(.2,.83),(.8,.83)],'#A8B4BB',1.3)
    elif kind == 'rank':
        for i,(yy,ww) in enumerate(((.12,.61),(.42,.43),(.72,.27))):
            el(.02,yy,.17,.17,pale,color)
            f.text(x+.045*size,y+(yy-.005)*size,.13*size,.2*size,str(i+1),size*.12,color,'center')
            rr(.28,yy,ww,.17,pale,r=3)
        ln([(.98,.16),(.98,.86)],width=1.8,arrow=True)
    elif kind == 'filter':
        poly([(.08,.1),(.93,.1),(.64,.48),(.64,.82),(.41,.94),(.41,.48)],pale)
        ln([(.2,.29),(.8,.29)],width=1.2)
    elif kind == 'decision':
        el(.36,.02,.28,.25,pale)
        ln([(.5,.29),(.5,.48)])
        ln([(.16,.64),(.16,.48),(.85,.48),(.85,.64)])
        ln([(.5,.48),(.5,.64)])
        for i,a in enumerate((.05,.39,.73)):
            rr(a,.69,.22,.24,PALETTES[i+1][0],color,r=4)
    elif kind == 'parameters':
        rr(.08,.08,.84,.82,WHITE,r=5)
        rr(.08,.08,.84,.2,pale,r=0)
        for yy in (.28,.49,.7): ln([(.08,yy),(.92,yy)],'#ADBAC3',1.2)
        ln([(.4,.08),(.4,.9)],'#ADBAC3',1.2)
        for yy in (.39,.6,.8):
            ln([(.17,yy),(.31,yy)],width=2)
            ln([(.5,yy),(.81,yy)],width=2)
    elif kind == 'blueprint':
        rr(.1,.03,.8,.93,pale,r=5)
        for aa,bb in ((.19,.18),(.6,.18),(.4,.55)):
            rr(aa,bb,.2,.17,WHITE,r=2)
        ln([(.39,.265),(.59,.265)],width=1.7,arrow=True)
        ln([(.29,.36),(.29,.635),(.39,.635)],width=1.7,arrow=True)
        ln([(.7,.36),(.7,.635),(.61,.635)],width=1.7,arrow=True)
        ln([(.22,.82),(.78,.82)],width=1.2)
    elif kind == 'compile':
        rr(.04,.08,.91,.66,WHITE,r=7)
        ln([(.37,.25),(.22,.4),(.37,.55)],width=2.7)
        ln([(.63,.25),(.78,.4),(.63,.55)],width=2.7)
        ln([(.55,.24),(.46,.58)],width=2.4)
        ln([(.5,.75),(.5,.91),(.27,.91),(.74,.91)],width=2)
    elif kind == 'curve':
        ln([(.12,.08),(.12,.86),(.97,.86)],width=1.6)
        ln([(.14,.84),(.23,.55),(.32,.3),(.41,.25),(.5,.4),(.63,.6),(.8,.74),(.96,.79)],width=2.8)
        ln([(.14,.84),(.27,.66),(.4,.44),(.51,.43),(.67,.61),(.83,.73),(.96,.8)],'#AA8174',1.8,dashed=True)
    elif kind == 'organ':
        # Stylized whole-body silhouette and organ compartments, not anatomy data.
        el(.4,.01,.2,.18,'#F2F3F4','#A5B0B6')
        poly([(.36,.21),(.25,.25),(.14,.56),(.22,.59),(.34,.37),(.34,.61),(.3,.97),(.43,.97),(.5,.68),(.57,.97),(.7,.97),(.66,.61),(.66,.37),(.78,.59),(.86,.56),(.75,.25),(.64,.21)],'#F2F3F4','#A5B0B6')
        el(.36,.29,.115,.18,'#BFD5E7','#739AB5')
        el(.53,.29,.115,.18,'#BFD5E7','#739AB5')
        el(.47,.36,.10,.11,'#DCAAA2','#A56860')
        poly([(.36,.49),(.61,.47),(.64,.52),(.54,.57),(.39,.56)],'#D9C3A0','#A58C63')
        el(.365,.585,.065,.1,'#CCBCD9','#8B729E')
        el(.57,.585,.065,.1,'#CCBCD9','#8B729E')
        rr(.45,.61,.1,.14,'#C7D8BF','#729569',3)
        ln([(.495,.45),(.495,.485)],'#A56860',1.6)
    elif kind == 'gut':
        poly([(.13,.11),(.84,.11),(.93,.28),(.9,.85),(.74,.94),(.66,.78),(.74,.65),(.73,.31),(.28,.31),(.28,.73),(.4,.82),(.34,.94),(.14,.85),(.08,.3)],'#EBC8B9','#A98171')
        ln([(.38,.39),(.62,.39),(.62,.49),(.38,.49),(.38,.59),(.62,.59),(.62,.69),(.46,.69)],'#A98171',2)
    elif kind == 'blood':
        rr(.07,.24,.87,.52,'#F5E8E5','#AE827B',11)
        for a,b in ((.19,.36),(.44,.43),(.69,.34)):
            el(a,b,.18,.15,'#D9A59E','#AE827B')
        ln([(.1,.13),(.84,.13)],'#AE827B',1.5,True)
    elif kind == 'liver':
        poly([(.07,.29),(.28,.13),(.72,.12),(.93,.24),(.89,.45),(.65,.63),(.47,.88),(.24,.84),(.13,.65)],'#DBB39D','#A88066')
        ln([(.63,.22),(.55,.51),(.48,.73)],'#A88066',1.4)
    elif kind == 'kidney':
        poly([(.29,.08),(.13,.17),(.06,.4),(.09,.66),(.24,.85),(.4,.78),(.45,.6),(.34,.46),(.45,.31),(.41,.14)],'#CFC1DE','#8D79A3')
        poly([(.71,.08),(.87,.17),(.94,.4),(.91,.66),(.76,.85),(.6,.78),(.55,.6),(.66,.46),(.55,.31),(.59,.14)],'#CFC1DE','#8D79A3')
        ln([(.35,.59),(.43,.79),(.44,.99)],'#8D79A3',1.5)
        ln([(.65,.59),(.57,.79),(.56,.99)],'#8D79A3',1.5)
    else:
        raise ValueError(kind)


def band(f,x,y,w,text,index,size=21):
    pale,border,color=PALETTES[index]
    f.rect(x,y,w,43,WHITE,border,8,1.1)
    f.text(x+15,y+7,w-30,30,text,size,color,'center')


def tile(f, cx, y, title, lines, kind, index, w=265, icon_size=72, size=25):
    pale,border,color=PALETTES[index]
    if kind=='adme':
        for j,(organ,label) in enumerate(zip(('gut','blood','liver','kidney'),('吸收','分布','代谢','排泄'))):
            xx=cx-115+j*64
            icon(f,organ,xx,y+1,39,color,pale)
            f.text(xx-5,y+48,49,22,label,15,color,'center')
            if j<3:f.line([(xx+42,y+23),(xx+58,y+23)],color,1.4,True)
    elif kind=='physiology':
        # Schematic blood/organ connections; labels and edges remain native.
        nodes={'血液':(cx,y+33),'肠道':(cx-89,y+9),'肝脏':(cx-89,y+63),
               '肾脏':(cx+89,y+9),'组织':(cx+89,y+63)}
        f.line([(cx-24,y+31),(cx-65,y+11)],color,1.4,True)
        f.line([(cx-89,y+23),(cx-89,y+49)],color,1.4,True)
        f.line([(cx-65,y+61),(cx-24,y+41)],color,1.4,True)
        f.line([(cx+24,y+22),(cx+65,y+2)],color,1.4,True)
        f.line([(cx+65,y+17),(cx+24,y+31)],color,1.4,True)
        f.line([(cx+24,y+38),(cx+65,y+55)],color,1.4,True)
        f.line([(cx+65,y+70),(cx+24,y+46)],color,1.4,True)
        for label,(xx,yy) in nodes.items():
            f.rect(xx-24,yy-13,48,27,WHITE,border,5,1.1)
            f.text(xx-22,yy-10,44,23,label,15,color,'center')
    else:
        icon(f,kind,cx-icon_size/2,y,icon_size,color,pale)
    f.text(cx-w/2,y+icon_size+14,w,39,title,size,INK,'center',True)
    f.text(cx-w/2,y+icon_size+58,w,62,'\n'.join(lines),20,MUTED,'center')


def quadrants():
    f=Figure('01_four_panel','四宫格 · 论文方法总图')
    f.heading('方案 01 / 四宫格')
    f.panel(48,184,880,353,'a','资料收集与建立知识库',0)
    f.panel(992,184,880,353,'b','组件推荐',1)
    f.panel(48,631,880,371,'c','模型推荐',2)
    f.panel(992,631,880,371,'d','蓝图设计',3)
    specs=[
        (48,184,0,[('多源资料收集',['文献 · 标签 · 临床 PK','公开模型与版本'],'documents'),
                   ('证据抽取与质控',['规则 + LLM 辅助','身份、单位与冲突校验'],'extract'),
                   ('知识图谱与索引',['药物—机制—器官—参数','证据库 + 模型目录'],'graph')],
         '可溯源证据、模型资产与版本化模板'),
        (992,184,1,[('建模需求解析',['药物 · 人群 · 给药方案','预测目标与约束'],'query'),
                    ('证据检索与重排',['关键词 + 向量 + 子图','相关性与机制一致性'],'retrieval'),
                    ('组件推断与校验',['组件、机制与参数候选','对应证据引用与约束'],'adme')],
         '输出：组件建议 + 结构化需求画像'),
        (48,631,2,[('候选模型召回',['具体模型版本与结构','场景、生物学与工程约束'],'model'),
                   ('需求匹配与排序',['机制覆盖与结构匹配','适用性与证据支持'],'rank'),
                   ('用途判断与选择',['复用 / 受控适配 / 拒绝','候选比较与推荐依据'],'decision')],
         '可靠模型 → 复用或受控适配 → 仿真与验证'),
        (992,631,3,[('结构模板与实例化',['无可靠复用模型时启动','房室、器官与过程连接'],'physiology'),
                    ('参数绑定与溯源',['组件需求 + 参数证据','量值、单位与来源'],'parameters'),
                    ('校验、编译与导出',['结构 · 量纲 · 质量守恒','版本化蓝图与仿真方案'],'compile')],
         '蓝图 + 参数清单 + 执行接口 → 仿真与验证'),
    ]
    for x,y,i,items,out in specs:
        f.label_context='abcd'[i]+' detail'
        centers=[x+150,x+440,x+730]
        for j,(title,lines,kind) in enumerate(items):
            tile(f,centers[j],y+87,title,lines,kind,i)
            if j<2:
                start=centers[j]+(135 if kind=='physiology' else 95)
                end=centers[j+1]-(135 if items[j+1][2]=='adme' else 95)
                f.line([(start,y+123),(end,y+123)],PALETTES[i][2],1.8,True)
        band(f,x+22,y+(288 if i<2 else 306),836,out,i)
    f.line([(937,359),(984,359)],PALETTES[0][2],2.5,True)
    f.line([(1432,548),(1432,576),(486,576),(486,620)],PALETTES[1][2],2.3,True)
    f.rect(642,552,632,45,WHITE,None,7)
    f.text(657,558,602,32,'需求画像用于模型匹配；组件与参数需求用于蓝图',21,MUTED,'center')
    f.line([(938,812),(981,812)],PALETTES[2][2],2.5,True)
    f.footer()
    return f


def swimlanes():
    f=Figure('02_four_lanes','四行流程 · 全链路展开')
    f.heading('方案 02 / 四行流程')
    rows=[
        ('a','资料收集与\n建立知识库',[
            ('多源资料汇集','文献、标签、PK 数据与模型','documents'),
            ('规则与 LLM 抽取','实体、机制、参数与证据位置','extract'),
            ('标准化与质量控制','身份归一、单位、去重与冲突','quality'),
            ('图谱与检索索引','可溯源证据 + 模型目录','graph')]),
        ('b','组件推荐',[
            ('用户需求解析','药物、人群、给药与预测目标','query'),
            ('多路检索与子图重排','关键词、语义与机制相关性','retrieval'),
            ('机制与组件推断','吸收、分布、代谢与排泄','components'),
            ('引用与约束校验','组件建议 + 结构化需求画像','quality')]),
        ('c','模型推荐',[
            ('候选模型召回','版本、结构与适用场景','model'),
            ('生物学与工程约束','药物、途径、人群与模型机制','filter'),
            ('需求匹配与排序','组件覆盖、任务匹配与证据','rank'),
            ('用途判断与模型选择','可靠模型 → 复用 / 受控适配','decision')]),
        ('d','蓝图设计',[
            ('结构模板选择','无可靠复用模型时启动','blueprint'),
            ('房室与过程实例化','器官、血流与 ADME 过程','organ'),
            ('有来源参数绑定','量值、单位、证据与不确定性','parameters'),
            ('校验、编译与导出','版本化蓝图 + 仿真方案','compile')]),
    ]
    for i,(letter,title,steps) in enumerate(rows):
        y=179+i*211
        pale,border,color=PALETTES[i]
        f.label_context=letter+' '+title.replace('\n','')
        f.rect(48,y,1824,190,pale,border,16,1.6)
        f.line([(287,y+16),(287,y+174)],border,1.3)
        f.text(70,y+19,198,44,f'({letter})',29,color,bold=True)
        f.text(70,y+70,200,91,title,29,color,bold=True)
        for j,(name,desc,kind) in enumerate(steps):
            cx=480+390*j
            icon(f,kind,cx-31,y+14,61,color,pale)
            f.text(cx-180,y+89,360,39,name,25,INK,'center',True)
            f.text(cx-185,y+137,370,30,desc,20,MUTED,'center')
            if j<3:f.line([(cx+89,y+49),(cx+300,y+49)],color,1.8,True)
        if i<3:
            # Inter-stage line stays in the gutter, never across a label.
            f.line([(1839,y+192),(1839,y+209)],color,1.8,True)
    f.footer()
    return f


def hub():
    f=Figure('03_central_framework','中心辐射 · 图形摘要')
    f.heading('方案 03 / 中心辐射')
    # Clockwise order: a (upper left), b (upper right), c (lower right), d (lower left).
    panels=[(48,181,0,'a','资料收集与建立知识库'),(992,181,1,'b','组件推荐'),
            (992,617,2,'c','模型推荐'),(48,617,3,'d','蓝图设计')]
    for x,y,i,l,t in panels:
        if l=='c':
            pale,border,color=PALETTES[i]
            f.label_context=l+' '+t
            f.rect(x,y,880,385,pale,border,19,1.8)
            f.text(x+220,y+17,50,43,f'({l})',29,color,bold=True)
            f.text(x+282,y+17,565,43,t,29,INK,bold=True)
            f.line([(x+220,y+69),(x+857,y+69)],border,1)
        else:
            f.panel(x,y,880,385,l,t,i)
    # Keep the central motif in a white clearance zone; all content stays in wings.
    data=[
        (48,181,0,'documents','graph',[
            ('多源资料','文献、标签、临床 PK、公开模型'),
            ('抽取与质控','规则 + LLM；身份、单位与溯源'),
            ('知识库构建','证据图谱、检索索引与模型目录')]),
        (992,181,1,'retrieval','components',[
            ('建模需求','药物、人群、给药与预测目标'),
            ('检索与推断','多路检索、子图重排与机制推断'),
            ('组件建议','ADME 组件、参数、引用与约束')]),
        (992,617,2,'model','rank',[
            ('候选与约束','模型版本、生物学与工程条件'),
            ('匹配与排序','需求画像、组件覆盖与证据支持'),
            ('模型选择','可靠候选 → 复用 / 受控适配')]),
        (48,617,3,'blueprint','parameters',[
            ('结构设计','模板选择、房室与过程实例化'),
            ('参数绑定','组件需求、量值、单位与来源'),
            ('蓝图输出','守恒校验、编译与版本化导出')]),
    ]
    for x,y,i,k1,k2,steps in data:
        pale,border,color=PALETTES[i]
        f.label_context='abcd'[i]+' detail'
        # Left panels keep text toward outer edge; right panels mirror it.
        xx=x+28 if x<500 else x+220
        for j,(title,desc) in enumerate(steps):
            yy=y+100+j*84
            f.ellipse(xx,yy+4,27,27,WHITE,border,1.3)
            f.text(xx+2,yy+4,24,24,str(j+1),16,color,'center')
            f.text(xx+42,yy-2,574,37,title,25,INK,bold=True)
            f.text(xx+42,yy+36,574,30,desc,21,MUTED)
        ix=x+701 if x<500 else x+28
        # Both icons stay above/below the central circle, never behind it.
        iy=290 if y<400 else 835
        icon(f,k1,ix,iy,70,color,pale)
        icon(f,k2,ix+88,iy,70,color,pale)
    # Continuous route arrows in the outer inter-panel gutters.
    f.line([(938,266),(982,266)],PALETTES[0][2],2.4,True)
    f.line([(1749,574),(1749,607)],PALETTES[1][2],2.4,True)
    f.line([(982,916),(938,916)],PALETTES[2][2],2.4,True)
    f.rect(745,934,430,36,WHITE,None,7)
    f.text(756,938,410,30,'无可靠复用模型 → 蓝图设计',21,PALETTES[2][2],'center')
    # Native organ diagram serves as the common scientific object, not a bitmap.
    f.ellipse(750,386,420,420,WHITE,'#D8E1E6',2.1)
    f.ellipse(771,407,378,378,'#FAFCFD','#E5EBEF',1.2)
    f.text(811,429,298,48,'PBPK',34,INK,'center',True)
    icon(f,'organ',875,489,170,'#7395A3')
    f.text(796,689,328,37,'证据驱动建模',25,INK,'center',True)
    f.text(796,733,328,30,'机制 · 结构 · 参数',20,MUTED,'center')
    # These labels describe the cross-module payloads without dense crossing lines.
    f.rect(563,569,282,37,WHITE,None,5)
    f.text(570,573,268,30,'证据 / 模型 / 模板',20,MUTED,'center')
    f.rect(1175,569,302,37,WHITE,None,5)
    f.text(1182,573,288,30,'组件建议 → 需求画像',20,MUTED,'center')
    f.footer()
    return f


def rgb(value): return RGBColor.from_string(value.lstrip('#'))
def px(value): return int(round(value*914400/120))


def export_ppt_slide(prs, figure):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.name=figure.title
    for item in figure.items:
        a=item.args
        if item.kind in ('rect','ellipse'):
            ty=MSO_SHAPE.OVAL if item.kind=='ellipse' else (MSO_SHAPE.ROUNDED_RECTANGLE if a['radius'] else MSO_SHAPE.RECTANGLE)
            sh=slide.shapes.add_shape(ty,px(a['x']),px(a['y']),px(a['w']),px(a['h']))
            if ty==MSO_SHAPE.ROUNDED_RECTANGLE:
                sh.adjustments[0]=min(.49,a['radius']/min(a['w'],a['h']))
            if a['fill']:sh.fill.solid();sh.fill.fore_color.rgb=rgb(a['fill'])
            else:sh.fill.background()
            if a['stroke']:sh.line.color.rgb=rgb(a['stroke']);sh.line.width=Pt(a['width']*.6)
            else:sh.line.fill.background()
        elif item.kind=='polygon':
            if a['arrowhead']:
                # PPT uses an attached native connector arrowhead instead.
                continue
            pts=a['pts']
            fb=slide.shapes.build_freeform(pts[0][0],pts[0][1],scale=914400/120)
            fb.add_line_segments(pts[1:],close=True)
            sh=fb.convert_to_shape()
            sh.fill.solid();sh.fill.fore_color.rgb=rgb(a['fill'])
            if a['stroke']:sh.line.color.rgb=rgb(a['stroke']);sh.line.width=Pt(a['width']*.6)
            else:sh.line.fill.background()
        elif item.kind=='line':
            for k,(p1,p2) in enumerate(zip(a['pts'],a['pts'][1:])):
                sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,px(p1[0]),px(p1[1]),px(p2[0]),px(p2[1]))
                sh.name=a['name']+f'/segment{k}'
                sh.line.color.rgb=rgb(a['color']);sh.line.width=Pt(a['width']*.6)
                if a['dashed']:sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
                if a['arrow'] and k==len(a['pts'])-2:
                    end=OxmlElement('a:tailEnd')
                    end.set('type','triangle');end.set('w','sm');end.set('len','sm')
                    sh.line._get_or_add_ln().append(end)
            continue
        elif item.kind=='text':
            sh=slide.shapes.add_textbox(px(a['x']),px(a['y']),px(a['w']),px(a['h']))
            tf=sh.text_frame
            tf.clear();tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE
            tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
            tf.vertical_anchor=MSO_ANCHOR.TOP
            for i,ln in enumerate(a['value'].split('\n')):
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
                p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[a['align']]
                p.space_before=p.space_after=Pt(0)
                p.line_spacing=Pt(a['size']*.6*1.28)
                run=p.add_run();run.text=ln
                run.font.name='Microsoft YaHei';run.font.size=Pt(a['size']*.6)
                run.font.bold=a['bold'];run.font.color.rgb=rgb(a['color'])
                ea=OxmlElement('a:ea');ea.set('typeface','Microsoft YaHei');run._r.get_or_add_rPr().append(ea)
        else:raise ValueError(item.kind)
        sh.name=a['name']+(('/'+a['value'].replace('\n',' ')[:44]) if item.kind=='text' else '')
    slide.notes_slide.notes_text_frame.text=(
        '预期完整研究框架的设计稿，不表示当前完成状态或实验结果。'
        '所有文字、图形、器官示意、框线、箭头均由 PowerPoint 原生可编辑对象组成，无栅格底图。'
        '模型选择区分复用、受控适配与拒绝；无可靠复用模型时进入蓝图设计。'
        '蓝图经仿真、校准与独立验证，不预设科学验证通过。'
        '参考用户附件的配色与科研排版；图像生成结果只用于设计探索。')


def export_images(f):
    root=ET.Element('{'+SVG_NS+'}svg',{'viewBox':f'0 0 {W} {H}','width':str(W),'height':str(H),'role':'img'})
    ET.SubElement(root,'{'+SVG_NS+'}title').text=f.title+' — PBPK 预期完整研究框架'
    im=Image.new('RGB',(W*SCALE,H*SCALE),WHITE)
    draw=ImageDraw.Draw(im)
    def el(tag,attrs):return ET.SubElement(root,'{'+SVG_NS+'}'+tag,{k:str(v) for k,v in attrs.items()})
    def bounds(a):return (a['x']*SCALE,a['y']*SCALE,(a['x']+a['w'])*SCALE,(a['y']+a['h'])*SCALE)
    for item in f.items:
        a=item.args
        if item.kind in ('rect','ellipse'):
            attrs={'fill':a['fill'] or 'none','stroke':a['stroke'] or 'none','stroke-width':a['width']}
            if item.kind=='rect':
                el('rect',dict(x=a['x'],y=a['y'],width=a['w'],height=a['h'],rx=a['radius'],**attrs))
                draw.rounded_rectangle(bounds(a),radius=a['radius']*SCALE,fill=a['fill'],outline=a['stroke'],width=max(1,round(a['width']*SCALE)))
            else:
                el('ellipse',dict(cx=a['x']+a['w']/2,cy=a['y']+a['h']/2,rx=a['w']/2,ry=a['h']/2,**attrs))
                draw.ellipse(bounds(a),fill=a['fill'],outline=a['stroke'],width=max(1,round(a['width']*SCALE)))
        elif item.kind=='polygon':
            el('polygon',{'points':' '.join(f'{x},{y}' for x,y in a['pts']),'fill':a['fill'],'stroke':a['stroke'] or 'none','stroke-width':a['width']})
            pp=[(x*SCALE,y*SCALE) for x,y in a['pts']]
            draw.polygon(pp,fill=a['fill'])
            if a['stroke']:draw.line(pp+[pp[0]],fill=a['stroke'],width=round(a['width']*SCALE),joint='curve')
        elif item.kind=='line':
            attrs={'points':' '.join(f'{x},{y}' for x,y in a['pts']),'fill':'none','stroke':a['color'],'stroke-width':a['width'],'stroke-linejoin':'round'}
            if a['dashed']:attrs['stroke-dasharray']='6 5'
            el('polyline',attrs)
            for (x1,y1),(x2,y2) in zip(a['pts'],a['pts'][1:]):
                length=((x2-x1)**2+(y2-y1)**2)**.5
                if not length:continue
                if a['dashed']:
                    for start in range(0,int(length),11):
                        end=min(start+6,length)
                        draw.line(((x1+(x2-x1)*start/length)*SCALE,(y1+(y2-y1)*start/length)*SCALE,
                                   (x1+(x2-x1)*end/length)*SCALE,(y1+(y2-y1)*end/length)*SCALE),fill=a['color'],width=max(1,round(a['width']*SCALE)))
                else:draw.line((x1*SCALE,y1*SCALE,x2*SCALE,y2*SCALE),fill=a['color'],width=max(1,round(a['width']*SCALE)))
        elif item.kind=='text':
            font=ImageFont.truetype(str(FONT),round(a['size']*SCALE))
            for i,line in enumerate(a['value'].split('\n')):
                y=a['y']+i*a['size']*1.28
                x=a['x']+({'left':0,'center':a['w']/2,'right':a['w']}[a['align']])
                attrs={'x':x,'y':y,'font-size':a['size'],'font-family':'Microsoft YaHei, Noto Sans CJK SC, PingFang SC, sans-serif',
                       'fill':a['color'],'dominant-baseline':'text-before-edge','text-anchor':{'left':'start','center':'middle','right':'end'}[a['align']]}
                if a['bold']:attrs['font-weight']='600'
                el('text',attrs).text=line
                draw.text((x*SCALE,y*SCALE),line,font=font,fill=a['color'],anchor={'left':'lt','center':'mt','right':'rt'}[a['align']],
                          stroke_width=1 if a['bold'] else 0)
    ET.ElementTree(root).write(OUT/(f.slug+'.svg'),encoding='utf-8',xml_declaration=True)
    im.save(OUT/(f.slug+'.png'),optimize=True)


def validate(prs_path,figures):
    ppt=Presentation(prs_path)
    assert len(ppt.slides)==3
    results=[]
    for slide,fig in zip(ppt.slides,figures):
        actual=[s.text for s in slide.shapes if s.has_text_frame and s.text]
        expected=[i.args['value'] for i in fig.items if i.kind=='text']
        assert actual==expected,(fig.slug,'text mismatch')
        for s in slide.shapes:
            assert s.left>=-2 and s.top>=-2 and s.left+s.width<=ppt.slide_width+2 and s.top+s.height<=ppt.slide_height+2,(fig.slug,s.name)
        svg=ET.parse(OUT/(fig.slug+'.svg'))
        assert not svg.findall('.//{'+SVG_NS+'}image')
        assert len(svg.findall('.//{'+SVG_NS+'}text'))==sum(len(t.split('\n')) for t in expected)
        results.append({'design':fig.slug,'native_shapes':len(slide.shapes),'editable_text_boxes':len(actual),'svg_text_elements':len(svg.findall('.//{'+SVG_NS+'}text'))})
    with ZipFile(prs_path) as z:
        assert not any(n.startswith('ppt/media/') for n in z.namelist()),'Raster/image asset embedded'
        for name in z.namelist():
            if name.endswith('.xml'):ET.fromstring(z.read(name))
    print(json.dumps({'pptx':str(prs_path),'checks':'all labels, canvas bounds, native shapes, no embedded images, valid XML','slides':results},ensure_ascii=False,indent=2))


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    figures=[quadrants(),swimlanes(),hub()]
    prs=Presentation()
    prs.slide_width=Inches(16);prs.slide_height=Inches(9)
    prs.core_properties.title='PBPK 建模推荐框架：三个可编辑设计方案'
    prs.core_properties.subject='资料与知识库、组件推荐、模型推荐、蓝图设计'
    prs.core_properties.author='PBPK research project'
    for fig in figures:
        export_images(fig)
        export_ppt_slide(prs,fig)
    target=OUT/'PBPK_workflow_editable.pptx'
    prs.save(target)
    validate(target,figures)


if __name__=='__main__':main()
