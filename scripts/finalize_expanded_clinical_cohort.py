#!/usr/bin/env python3
"""Audit the new curator packet without touching frozen inputs or predictions."""
from __future__ import annotations
import importlib.util
import json
import os
from pathlib import Path
import re
import xml.etree.ElementTree as ET
from docx import Document
from lxml import html

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('curator', ROOT / 'scripts/curate_expanded_clinical_cohort.py')
cur = importlib.util.module_from_spec(spec)
spec.loader.exec_module(cur)
OUT = cur.OUT


def load(ref):
    path = ROOT / ref['path']
    assert cur.ref(path)['sha256'] == ref['sha256'], f'Hash mismatch: {path}'
    return json.loads(path.read_text())


def document_text(path):
    if path.suffix == '.xml':
        return cur.text(ET.parse(path).getroot())
    if path.suffix == '.docx':
        doc = Document(path)
        return '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    if path.suffix == '.html':
        doc = html.fromstring(path.read_bytes())
        for node in doc.xpath('//script|//style'):
            node.drop_tree()
        return ' '.join(doc.itertext())
    return path.read_text()


def main():
    os.sched_setaffinity(0, sorted(os.sched_getaffinity(0))[:2])
    receipt_path = OUT / 'curation_receipt.json'
    receipt = json.loads(receipt_path.read_text())
    aliases, checks = [], []
    for row in receipt['records']:
        drug = row['compound']
        protocol, audit, sealed = [load(row[k]) for k in ('protocol', 'source_audit', 'sealed_endpoints')]
        source_spec = cur.specifications()[drug]
        meta = ET.parse(cur.article_path(source_spec['pmc'])).getroot().find('.//article-meta')
        first_author = cur.text(meta.find('.//contrib//surname'))
        years = {cur.text(p.find('year')) for p in meta.findall('pub-date')}
        if drug == 'pitavastatin':
            years.add('2023')  # 2022 online publication / 2023 volume16 issue1.
        registry_aliases = sorted({re.sub(r'^EudraCT', '', t, flags=re.I) for t in source_spec['trials']})
        hits = []
        for doc in audit['audit_documents']:
            path = ROOT / doc['path']
            assert cur.ref(path)['sha256'] == doc['sha256']
            text = document_text(path)
            matched = [t for t in registry_aliases if cur.norm(t) in cur.norm(text)]
            for year in years:
                matched += re.findall(rf'\b{re.escape(first_author)}\b[^\n.{{}}]{{0,45}}\b{year}\b', text, re.I)
            if matched:
                hits.append({'path': doc['path'], 'matched_aliases': matched})
        assert not hits, f'Clinical identity alias needs review: {drug}'
        aliases.append({'compound': drug, 'first_author': first_author,
                        'publication_year_aliases': sorted(years), 'bare_trial_identifiers_checked': registry_aliases,
                        'audit_documents_checked': len(audit['audit_documents']), 'matches': [],
                        'supplements_and_abbreviated_author_year_rechecked': True,
                        'qualification_scope_unchanged': audit['source_qualification_scope']})
        assert not {'value', 'source_cell', 'comparable_value'}.intersection(set().union(*(set(e) for e in protocol['endpoint_contracts'])))
        assert protocol['run_ready'] is False and protocol['actual_evaluation_completed'] is False
        assert protocol['case_id'] == sealed['case_id'] == row['case_id']
        assert len(sealed['endpoints']) == row['endpoint_count']
        assert sum(e['evaluable_from_source'] for e in sealed['endpoints']) == row['source_evaluable_endpoint_count']
        for e in sealed['endpoints']:
            assert cur.scalar(e['source_cell']) == e['value']
            assert e['value'] * e['conversion_multiplier'] == e['comparable_value']
        checks.append({'case_id': row['case_id'], 'hashes_verified': True,
                       'endpoint_cells_and_unit_conversions_verified': True,
                       'observed_endpoint_values_absent_from_protocol': True,
                       'sealed_file_mode_0600': oct((ROOT / row['sealed_endpoints']['path']).stat().st_mode & 0o777) == '0o600'})
    by_drug = {r['compound']: load(r['protocol']) for r in receipt['records']}
    assert by_drug['ethinylestradiol']['cohort_id'] == by_drug['levonorgestrel']['cohort_id']
    assert len({r['compound'] for r in receipt['records']}) == 9
    alias_ref = cur.save(OUT / 'identity_alias_recheck.json', {'status': 'passed', 'records': aliases})
    validation_ref = cur.save(OUT / 'final_validation.json', {'status': 'passed', 'checks': checks,
        'paired_contraceptive_compounds_clustered_as_one_cohort': True, 'prediction_or_score_files_read': False})

    supplement_notes = {
        'dapagliflozin': 'Main article tables contain NCA endpoints and BE statistics; no supplemental numerical time-course link identified.',
        'digoxin': 'Official DOCX supplementary file inspected: assay validation; no exact clinical concentration-time table.',
        'omeprazole': 'Seven main article tables contain demographics/endpoints/BE/safety, with no supplementary attachment linked in full XML.',
        'pitavastatin': 'All five official DOCX tables inspected: demographics, dissolution, ANOVA, preliminary and confirmatory NCA; no complete human concentration-time table.',
        'quetiapine': 'Official nested-ZIP supplementary PDF,2pages: genotype/phenotype NCA summary tables only, no individual weight linkage or complete time course.',
        'vancomycin': 'Official nested-ZIP supplementary PDF,6pages: model parameters/covariate selection and individual fit/VPC figures; no exact numerical concentration-time table. Figures were not digitized.',
        'bedaquiline': 'Table5 has two selected incurred-sample-reanalysis measurements per person; peak/tail-selected sparse assay checks are not complete curves and are not counted as full independent curves.',
        'ethinylestradiol': 'Previously downloaded official31kB DOCX inspected: eligibility/bioanalysis/semaglutide-SNAC endpoint tables; no exact EE/LNG concentration-time grid or values.',
        'levonorgestrel': 'Shares the preceding Trial1 publication and official DOCX; no separate curve or trial counted.'}
    supplements = [p for p in (OUT / 'sealed/articles').glob('*_supplements.zip')]
    old_supp = cur.OLD / 'protocol_revision_20260920/curator_only/PMC8416862_supplement_europepmc.docx'
    availability_ref = cur.save(OUT / 'curve_availability_audit.json', {
        'scope': 'Selected nine drugs: full primary XML tables, explicit official attachments, and existing DOCX; no image digitization or inferred curves.',
        'records': [{'compound': d, 'finding': n, 'exact_full_numeric_curve_ready': False} for d, n in supplement_notes.items()],
        'inspected_supplement_archives': [cur.ref(p) for p in sorted(supplements) if p.stem.split('_')[0] in {'PMC9841298','PMC12925031','PMC8540141','PMC11504268'}],
        'reused_previously_downloaded_supplement': cur.ref(old_supp),
        'new_full_curve_drugs': 0, 'existing_curve_drugs_unchanged': ['voriconazole', 'atorvastatin', 'rosuvastatin']})
    rejections = [
        ('isoniazid', 'PMC9343941', 'Pooled older Chen2009 NAT2 cohort with other BE/patient data; source-disjoint observed endpoints cannot yet be isolated.'),
        ('rifampicin', 'PMC12795710', 'Assay-validation report, no actual new human PK endpoint cohort identified.'),
        ('rifampicin', 'PMC11303472', 'Actual healthy study but article/supplement provide model estimates and figures, not an exact clinical curve or NCA endpoint table extracted for validation.'),
        ('ruxolitinib', 'PMC12030279', 'Beagle/rat experiment; not a human clinical validation cohort.'),
        ('probenecid', 'PMC11271292', 'Dose-history inconsistency: five prior doses stated but six administration times listed. Two sparse sampling times are not a complete curve; do not guess regimen.'),
    ]
    cur.save(OUT / 'candidate_selection_audit.json', {'counted_drugs': sorted(by_drug),
        'excluded_or_unresolved_not_counted': [{'compound': d, 'pmcid': p, 'reason': r} for d, p, r in rejections],
        'selection_used_no_model_predictions_or_scores': True,
        'other_query_hits_not_automatically_qualified': True})
    markers = [json.loads(p.read_text()) for p in OUT.glob('sealed/**/*.acquisition.json')]
    transfer = sum(sum(a.get('bytes', a.get('bytes_received', 0)) for a in e.get('attempts', [])) for e in markers)
    assert transfer < 80 * 1024**2 and cur.cohort_bytes() < 100 * 1024**2
    budget_ref = cur.save(OUT / 'bounded_acquisition_receipt.json', {
        'completed_requests': len(markers), 'new_transfer_bytes_including_failed_received_bytes': transfer,
        'current_cumulative_curation_bytes_before_receipt': cur.cohort_bytes(),
        'maximum_new_transfer_bytes': 80 * 1024**2, 'maximum_cumulative_curation_bytes': 100 * 1024**2,
        'maximum_document_bytes': 10 * 1024**2, 'maximum_parallel_downloads': 3,
        'maximum_request_seconds': 30, 'maximum_retries': 1,
        'failed_requests_not_automatically_repeated': [e for e in markers if e['status'] != 'downloaded'],
        'budget_exhausted': False, 'downloads_running': False})
    old = json.loads((cur.OLD / 'acquisition_summary.json').read_text())
    drugs = sorted(set(old['curation_coverage']['drug_names']) | set(by_drug))
    assert len(drugs) == 30
    handoff = dict(schema_version='pbpk-new-clinical-cohort-handoff-v1', status='curation_and_static_verification_complete',
                   curation_receipt=cur.ref(receipt_path), identity_alias_audit=alias_ref,
                   final_static_validation=validation_ref, curve_availability_audit=availability_ref,
                   acquisition_receipt=budget_ref, parent_baseline_summary=cur.ref(cur.OLD / 'acquisition_summary.json'),
                   new_drugs=sorted(by_drug), new_drug_count=9, candidate_drugs_total=30, drug_names=drugs,
                   clinical_protocol_metadata_total=32, unique_target_publications_total=27,
                   new_target_cohorts=8, new_endpoint_records=24, source_evaluable_endpoint_records=23,
                   new_full_numeric_curve_drugs=0, full_curve_drugs_total=3,
                   new_cases_native_verified=0, new_cases_evaluated=0, target_independent_evaluation_drugs=60,
                   remaining_candidate_gap_to_60=30, clinical_metadata_count_is_not_evaluation_count=True,
                   old_frozen_protocols_modified=False, predictions_or_scores_read=False,
                   qualifying_protocols_directory=str((OUT/'protocols').relative_to(ROOT)))
    final = cur.save(OUT / 'handoff_manifest.json', handoff)
    print(json.dumps({'manifest': final, 'candidate_drugs_total': 30, 'new_drug_count': 9,
                      'new_evaluated_cases': 0, 'new_transfer_bytes': transfer,
                      'cumulative_curation_bytes': cur.cohort_bytes()}, ensure_ascii=False))


if __name__ == '__main__':
    main()
