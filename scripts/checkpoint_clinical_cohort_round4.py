#!/usr/bin/env python3
"""Read-only source validation and immutable checkpoint; no network or predictions."""
from pathlib import Path
import hashlib
import json
import os
import re
import unicodedata
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from lxml import html

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'experiments/blueprint_design/external_validation/clinical_cohort_round4_20260920'
OLD = ROOT / 'experiments/blueprint_design/independent_cohort_20260920'
PARENT = ROOT / 'experiments/blueprint_design/external_validation/expanded_cohort_20260920/handoff_manifest.json'
DIRS = [OLD] + [ROOT / 'experiments/blueprint_design/external_validation' / n for n in (
    'newstudy_bundle_20260920', 'additional_native_cohort_20260920', 'concentration_curve_audit_20260920',
    'remaining_cohort_20260920', 'expanded_cohort_20260920', 'clinical_cohort_round4_20260920')]


def text(node):
    return ' '.join(''.join(node.itertext()).split()) if node is not None else ''


def normalize(s):
    return re.sub('[^a-z0-9]', '', unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode().lower())


def ref(path):
    return dict(path=str(path.relative_to(ROOT)), sha256=hashlib.sha256(path.read_bytes()).hexdigest(), bytes=path.stat().st_size)


def footprint():
    return sum(p.stat().st_size for d in DIRS for p in d.rglob('*') if p.is_file())


def save(name, value):
    path = BASE / name
    blob = (json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n').encode()
    if path.exists():
        assert path.read_bytes() == blob, 'Do not overwrite an earlier checkpoint'
    else:
        assert footprint() + len(blob) < 100 * 1024**2
        path.write_bytes(blob)
    return ref(path)


def cases():
    return [
        ('esomeprazole','PMC10175790',['CTR20171347','CTR20171484'], 'Healthy Chinese; fasting reference Nexium20mg;32 people.',
         ['Original Kanacher2020 paper/full-reference audit must be attached.', 'Preserve author-only corrigendum.', 'Tablet release differs from source capsule; requires explicit mapping.']),
        ('dabigatran','PMC11909909',['CTR20231968'], 'Healthy Chinese; fasting150mg dabigatran etexilate reference;50 enrolled; fully replicated crossover.',
         ['Audit Moj2019 original model clinical sources.', 'Define total/free assay pool and conjugate handling; do not equate free with protein-unbound.', 'Confirm endpoint-evaluable participants versus repeated reference exposures.']),
        ('posaconazole','PMC10986696',[], 'Healthy adults18–55;100mg reference gastro-resistant tablet fasting;36 randomized/33 paired comparison.',
         ['Exact nominal sampling grid absent;144h horizon known.', 'Keep first bioequivalence subset separate from food-effect paired subset.', 'Finish original model clinical-study/source audit.']),
        ('efavirenz','PMC6373243',[], 'Healthy subjects; efavirenz-alone day1 cohort20 before rolapitant;17men/3women.',
         ['Extract dose/formulation, meal and sampling protocol.', 'Attach original source model clinical report audit.', 'Publication already used for another probe; deduplicate global publication counts.']),
        ('risperidone','PMC3627046',[], 'Healthy Chinese24men;2mg Risperdal reference tablet; crossover.',
         ['Attach complete source audit including Kneller full clinical supplement and Rudesheim2022 clinical-source article.', 'Source formulation Solution does not establish tablet equivalence.']),
        ('ruxolitinib','PMC6221065',['INCB18424-139'], 'Healthy phase1 cohort9;IR25mg period1;6men/3women; exclude myelofibrosis phase2.',
         ['Attach original Gerner2022 full source audit.', 'Do not assign phase2 NCT01340651 to healthy phase1.', 'Nominal grid exists in downloaded DOCX supplement.']),
        ('olanzapine','PMC3480756',[], 'Healthy18men; reference suspension10mg/5mL versus testODT10mg.',
         ['Source NCA table is image-only; no complete table extraction saved.', 'Audit original Zheng2022 clinical sources.', 'Do not interchange test/reference arms.']),
        ('cabozantinib','PMC12960307',['NCT05263245'], 'Metastatic renal-cell carcinoma12patients; mixed20/40/60mg once-daily doses; not healthy.',
         ['Audit original Gerner2021 paper.', 'Raw pooled endpoints are not dose-normalized; individual dose/covariate linkage unavailable.', 'Absolute mean aggregation definition unresolved.', 'Supplement contains menus/power calculations, no individual concentration table.']),
        ('raltegravir','PMC8325523',['NCT03205566'], 'HIV-negative36participants;400mg twice-daily; staggered sampling subgroups.',
         ['Main text and official registry say7days while figure caption says14; preserve discrepancy.', 'No ordinary NCA endpoint table extracted; registry protection thresholds are pharmacodynamic, not PK endpoints.', 'Whole supplemental ZIP exceeded10MiB and was not saved; exact attachment/protocol not retrieved.', 'Staggered dose/sampling/period linkage remains incomplete.']),
    ]


def source_path(pmc):
    if pmc == 'PMC6373243':
        return OLD / 'retrieval_round1/sealed_articles/PMC6373243.xml'
    return BASE / 'sealed/articles' / (pmc + '.xml')


def source_content(path):
    if path.suffix == '.xml':
        root = ET.parse(path).getroot()
        return text(root), len(list(root.iter('ref')))
    if path.suffix == '.docx':
        doc = Document(path)
        return '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells]), None
    if path.suffix == '.html':
        root = html.fromstring(path.read_bytes())
        for node in root.xpath('//script|//style'):
            node.drop_tree()
        return ' '.join(root.itertext()), len(root.xpath('//li[contains(@class,"c-article-references__item")]'))
    return path.read_text(), None


def main():
    os.sched_setaffinity(0, sorted(os.sched_getaffinity(0))[:2])
    parent = json.loads(PARENT.read_text())
    inventory = json.loads((OLD / 'cohort_inventory.json').read_text())
    checked, failures = [], []
    transfer = 0
    for marker in sorted(BASE.rglob('*.acquisition.json')):
        item = json.loads(marker.read_text())
        transfer += sum(a.get('bytes_received', 0) for a in item.get('attempts', []))
        path = ROOT / item['path']
        if item['status'] == 'downloaded':
            assert path.is_file()
            actual = ref(path)
            assert actual['sha256'] == item['sha256'] and actual['bytes'] == item['bytes']
            if path.suffix == '.xml': ET.parse(path)
            if path.suffix == '.json': json.loads(path.read_text())
            if path.suffix in {'.zip', '.docx'}:
                with zipfile.ZipFile(path) as z:
                    assert z.testzip() is None
            checked.append(dict(file=actual, acquisition_record=ref(marker), parsed_or_container_checked=path.suffix in {'.xml','.json','.zip','.docx'}))
        else:
            assert not path.exists(), 'Failed acquisition unexpectedly has payload'
            failures.append(dict(acquisition_record=ref(marker), url=item['url'], reason=item.get('reason'), attempts=len(item.get('attempts', [])), failed_payload_absent=True))
    records=[]
    for drug,pmc,trials,population,pending in cases():
        path=source_path(pmc); root=ET.parse(path).getroot(); meta=root.find('.//article-meta')
        ids={x.get('pub-id-type'):text(x) for x in meta.findall('article-id') if x.get('pub-id-type') in {'doi','pmid','pmcid'}}
        assert ids.get('pmcid') == pmc and ids.get('doi')
        title=text(meta.find('./title-group/article-title'))
        dates=[dict(attributes=x.attrib, components={c.tag:text(c) for c in x}) for x in meta.findall('pub-date')]
        authors=[dict(surname=text(a.find('surname')),given_names=text(a.find('given-names'))) for a in meta.findall('./contrib-group/contrib/name')]
        first=authors[0]['surname'];years=sorted({text(y) for d in meta.findall('pub-date') for y in d.findall('year')})
        docs={}
        alias='posaconazol' if drug=='posaconazole' else drug
        for row in inventory['drug_inventory']:
            if row['compound']==alias:
                for model in row['models']:
                    p=ROOT/model['snapshot_path'];docs[p]='registered_same_compound_snapshot'
                    docs.update({d:'repository_source_document' for d in p.parent.rglob('*.md')})
        extra={'risperidone':[BASE/'sealed/model_papers/kneller2019.html',BASE/'sealed/model_papers/kneller2019_supplement.docx'],
               'ruxolitinib':[BASE/'sealed/articles/PMC9785192.xml'],'posaconazole':[BASE/'sealed/articles/PMC9785192.xml'],
               'olanzapine':[BASE/'sealed/articles/PMC8807508.xml']}.get(drug,[])
        docs.update({p:'original_model_document' for p in extra})
        tokens=list(ids.values())+[title]+trials
        audits=[]
        for p,kind in sorted(docs.items()):
            content,nrefs=source_content(p);normalized=normalize(content)
            matches=[t for t in tokens if normalize(t) in normalized]
            author_year=[f'{first}{year}' for year in years if normalize(first+year) in normalized]
            audits.append(dict(document=ref(p),kind=kind,target_identifier_or_title_matches=matches,first_author_year_flags=author_year,references_present=nrefs))
        records.append(dict(compound=drug,source=ref(path),bibliographic_metadata=dict(identifiers=ids,title=title,authors=authors,publication_dates_as_encoded=dates),
            clinical_population_and_arm=population,trial_or_study_ids=trials,source_audit_screen=dict(documents=audits,complete_independence_qualification=False,
            exact_identity_matches=sum(len(a['target_identifier_or_title_matches']) for a in audits),screen_not_a_substitute_for_clinical_source_lineage_review=True),
            stage='raw_primary_source_saved_local_curation_pending',normalized_safe_protocol_complete=False,normalized_sealed_endpoints_complete=False,
            native_verified=False,evaluated=False,pending=pending))
    supplement=BASE/'safe_protocol_supplements/clarithromycin_oh2023_dose_history_population.json'
    receipt=dict(schema_version='pbpk-source-acquisition-checkpoint-v1',status='saved_sources_validated_partial_curation',
        source_candidates=records,new_primary_source_candidate_drugs=len(records),new_completed_safe_protocols=0,new_completed_sealed_endpoint_packets=0,
        completed_existing_drug_safe_supplements=[ref(supplement)],download_verification=dict(successful_files=len(checked),failed_documents=len(failures),success_records=checked,failures=failures,
        recorded_transferred_bytes_lower_bound=transfer,oversize_stream_additional_bytes_upper_bound=65536,
        oversize_reason='One Raltegravir supplemental ZIP reached10MiB; next chunk triggered stop, no payload was written. Recorded transfer excludes at most one65536-byte chunk.',
        cohort_bytes_at_checkpoint=footprint(),cumulative_limit_bytes=100*1024**2,no_new_network_calls_in_checkpoint=True),
        no_prediction_or_score_files_read=True,no_solver_or_gpu_used=True,old_frozen_files_modified=False,
        prior_completed_handoff=ref(PARENT),script=ref(Path(__file__)))
    rref=save('source_receipt.json',receipt)
    handoff=dict(schema_version='pbpk-new-clinical-cohort-handoff-v1',status='saved_sources_checkpoint_no_new_qualified_protocols',source_receipt=rref,parent_baseline_summary=ref(PARENT),
        new_drugs=[],new_drug_count=0,candidate_drugs_total=parent['candidate_drugs_total'],drug_names=parent['drug_names'],
        clinical_protocol_metadata_total=parent['clinical_protocol_metadata_total'],unique_target_publications_total=parent['unique_target_publications_total'],
        new_target_cohorts=0,new_endpoint_records=0,source_evaluable_endpoint_records=0,new_full_numeric_curve_drugs=0,
        full_curve_drugs_total=parent['full_curve_drugs_total'],new_cases_native_verified=0,new_cases_evaluated=0,
        target_independent_evaluation_drugs=60,remaining_candidate_gap_to_60=60-parent['candidate_drugs_total'],
        clinical_metadata_count_is_not_evaluation_count=True,old_frozen_protocols_modified=False,predictions_or_scores_read=False,
        pending_saved_source_candidate_drugs=[r['compound'] for r in records],pending_saved_source_candidate_count=len(records),
        completed_existing_drug_safe_supplements=[ref(supplement)],checkpoint_does_not_discard_saved_partial_work=True)
    href=save('handoff_manifest.json',handoff)
    print(json.dumps(dict(handoff=href,source_receipt=rref,successful_files=len(checked),failed_documents=len(failures),cohort_bytes_after_checkpoint=footprint(),candidate_drugs_total=handoff['candidate_drugs_total'],pending_source_candidates=len(records))))


if __name__ == '__main__':main()
