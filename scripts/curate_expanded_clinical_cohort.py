#!/usr/bin/env python3
"""Create nine source-disjoint clinical metadata/endpoint packets, without fitting.

All observed values and source tables are sealed. Source mistakes remain visible
in qualification metadata; no guessed corrections, digitization, or generated PK.
"""
from __future__ import annotations
import ast
import hashlib
import json
import os
from pathlib import Path
import re
import unicodedata
import xml.etree.ElementTree as ET
from docx import Document
from lxml import html

ROOT = Path(__file__).resolve().parents[1]
OLD = ROOT / 'experiments/blueprint_design/independent_cohort_20260920'
OUT = ROOT / 'experiments/blueprint_design/external_validation/expanded_cohort_20260920'
ASSETS = ROOT / 'artifacts/pbpk-model-assets/public_collection_20260909'
SOURCE = OUT / 'sealed/articles'
PAPERS = OUT / 'sealed/model_papers'


def ref(p):
    return {'path': str(p.relative_to(ROOT)), 'sha256': hashlib.sha256(p.read_bytes()).hexdigest(), 'bytes': p.stat().st_size}


def cohort_bytes():
    dirs = [OLD] + [ROOT / 'experiments/blueprint_design/external_validation' / n for n in ('newstudy_bundle_20260920', 'additional_native_cohort_20260920', 'concentration_curve_audit_20260920', 'remaining_cohort_20260920', 'expanded_cohort_20260920')]
    return sum(p.stat().st_size for d in dirs for p in d.rglob('*') if p.is_file())


def save(p, obj, sealed=False):
    data = (json.dumps(obj, ensure_ascii=False, indent=2, allow_nan=False) + '\n').encode()
    assert cohort_bytes() + len(data) < 100 * 1024**2, 'cumulative_cohort_100MiB_limit'
    if p.exists():
        assert p.read_bytes() == data, 'Refusing to overwrite changed output; package may be frozen'
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(data)
    if sealed:
        p.chmod(0o600)
        p.parent.chmod(0o700)
    return ref(p)


def text(e):
    return ' '.join(''.join(e.itertext()).split()) if e is not None else ''


def norm(s):
    return re.sub('[^a-z0-9]', '', unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode().lower())


def specifications():
    mean = 'arithmetic_mean_of_individual_NCA_endpoints'
    cases = {
        'dapagliflozin': dict(pmc='PMC10128151', snapshot='Dapagliflozin-Model__440e09902dde/source/Dapagliflozin-Model.json', compound='Dapagliflozin',
            suffix='nct04938752_forxiga10mg_reference', trials=['NCT04938752'], n=28,
            dose_mg=10, formulation='Forxiga dapagliflozin propanediol monohydrate reference tablet', route='oral', dose_times_h=[0],
            dose_basis='10mg labelled active dapagliflozin; not 10mg prodrug or propanediol salt mass',
            sampling_h=[0,5/60,10/60,.25,.5,.75,1,1.25,1.5,2,3,4,5,6,8,10,12,24,48], origin='single reference dose',
            food='Not explicitly specified by article or NCT04938752 registry; fasting glucose assessment is not proof of fasting dosing.',
            population=dict(healthy=True,ethnicity='Korean',sex='male',age_mean_year=24.20,height_mean_cm=175.91,weight_mean_kg=71.49,bmi_mean=23.10,demographic_n=30,pk_n=28,age_range_eligibility=[19,45]),
            table='t0001', endpoints=[('Cmax',2,2,'microgram/L','ng/mL',1),('AUC_last',3,2,'microgram*h/L','ng*h/mL',1),('AUC_infinity',4,2,'microgram*h/L','ng*h/mL',1)],
            aggregation=mean, papers=[], native_formulation='IC tablet (Chang 2015)',
            limitations=['Food state must be registered as unknown/reference assumption; do not infer it from fasting glucose.', 'Demographics describe30 randomized, not separately28 evaluable.', 'Use conventional dapagliflozin reference only; ester prodrug test is excluded.']),
        'digoxin': dict(pmc='PMC12925031',snapshot='Digoxin-Model__1550093a7102/source/Digoxin-Model.json',compound='Digoxin',
            suffix='ctr20221249_baseline05mg',trials=['CTR20221249'],n=16,
            dose_mg=.5,formulation='Two0.25mg Sanofi Digoxin tablets, lotBHG0746',route='oral',dose_times_h=[0],dose_basis='labelled digoxin active drug',
            sampling_h=[.25,.5,1,1.5,2,2.5,3,4,6,8,12,24,36,48,72,96,144,168],predose_interval_h=[-1,0],origin='day1 dose before any PB-119',
            food='At least10h overnight fast; first breakfast at least1h after dose; standard lunch4h after.',
            population=dict(healthy=True,ethnicity='Chinese',sex='male',n=16,age_mean_year=30.0,height_mean_cm=172.31,weight_mean_kg=65.96,bmi_mean=22.20),
            table='table3-20503121261424728',endpoints=[('AUC_infinity',1,1,'ng*h/mL','ng*h/mL',1),('AUC_last',2,1,'ng*h/mL','ng*h/mL',1),('Cmax',3,1,'ng/mL','ng/mL',1)],
            aggregation=mean,papers=['hanke2018'],native_formulation='Solution',
            limitations=['Exact predose sample time lies within1h before dose, not registered as an exact zero-hour draw.', 'Published source model uses Solution; tablet release equivalence must be checked or explicitly approximated.', 'Combination period and warfarin cohort excluded; all publication data denied to development.']),
        'omeprazole':dict(pmc='PMC13525694',snapshot='Omeprazole-Model__c8b0f70a40c8/source/Omeprazole-Model.json',compound='Omeprazole',
            suffix='fasted_losec20mg_reference',trials=['CTR20220627','CTR20221098'],n=48,
            dose_mg=20,formulation='Losec20mg enteric-coated capsule, AstraZeneca, lot06BN020',route='oral',dose_times_h=[0],dose_basis='labelled total racemic omeprazole dose',
            sampling_h=[-1,.25,.5,.75,1,1.33,1.67,2,2.33,2.67,3,3.5,4,5,6,8,10,12,14],origin='single fasting reference dose; actual predose nominal draw at-1h',
            food='At least10h overnight fasting, approximately240mL warm water.',
            population=dict(healthy=True,ethnicity='Chinese',n=48,sex_counts=dict(male=43,female=5),age_mean_year=26.70,height_mean_cm=168.04,weight_mean_kg=63.05,bmi_mean=22.28),
            table='Tab2',endpoints=[('Cmax',3,2,'ng/mL','ng/mL',1),('AUC_last',4,2,'ng*h/mL','ng*h/mL',1),('AUC_infinity',5,2,'ng*h/mL','ng*h/mL',1)],
            aggregation=mean,papers=['PMC7764797'],native_formulation='Omeprazole capsule',
            disabled_endpoints={'AUC_infinity':'Source reference AUCinfinity is below AUClast with the same reportedn48; original value retained without guessed correction.'},
            limitations=['Two trial IDs are publication-level; exact mapping of each ID to fasting/fed arm is not guessed.', 'Reference AUCinfinity is quarantined for internal source inconsistency; Cmax and AUClast remain separate traceable targets.', 'Total racemate assay must not be compared with a single-enantiomer prediction.', 'Predose-1h is not shifted to zero.']),
        'pitavastatin':dict(pmc='PMC9841298',snapshot='Pitavastatin-Model__84855836ef76/source/Pitavastatin-Model.json',compound='Pitavastatin',
            suffix='confirmatory_n36_reference2mg',trials=['CTR20190940','CTR20191166'],n=36,
            dose_mg=2,formulation='2mg reference pitavastatin calcium tablet',route='oral',dose_times_h=[0],dose_basis='2mg source-labelled pitavastatin dose; calcium formulation not reinterpreted as2mg salt mass',
            sampling_h=[0,.17,.33,.5,.67,1,1.25,1.5,2,2.5,3,3.5,4,5,6,8,12,24,36,48],origin='single reference dose in confirmatory36-person crossover',
            food='Source study states fasting; exact overnight duration not supplied in inspected text.',
            population=dict(healthy=True,n=36,sex_counts=dict(male=21,female=15),age_mean_year=(23.94+24.72)/2,weight_mean_kg=(60.77+60.94)/2,height_mean_cm=(168.10+164.40)/2,bmi_mean=(21.49+22.43)/2,demographic_aggregation='Derived from two equally sized18-subject randomization sequences; no pooled SD reconstructed.',ethnicity_counts=dict(Han=35,other=1)),
            docx='PMC9841298/CTS-16-85-s005.docx',table=0,endpoints=[('Cmax',1,2,'ng/mL','ng/mL',1),('AUC_last',2,2,'ng*h/mL','ng*h/mL',1),('AUC_infinity',3,2,'ng*h/mL','ng*h/mL',1)],
            aggregation=mean,papers=[],native_formulation='tablet',
            limitations=['Uses actual observed reference column in supplementalTableS4, not main-text PBPK predictions or virtual BE results.', 'Preliminary24-person study not mixed with36-person confirmatory study.', 'Entire paper and both CTR IDs excluded from all development; exact CTR-to-study mapping not invented.', 'Reference n36 retained separately from test n35.']),
        'quetiapine':dict(pmc='PMC8540141',snapshot='Quetiapine-Model__0564367791d8/source/Quetiapine.json',compound='Quetiapine',
            suffix='eudract202000109114_trial2_reference50mg',trials=['EudraCT2020-001091-14','EudraCT2018-003079-37'],n=30,
            dose_mg=50,formulation='Two25mg Seroquel film-coated reference tablets',route='oral',dose_times_h=[0],dose_basis='labelled quetiapine equivalent; not fumarate salt mass',
            sampling_h=[0,.33,.67,1,1.5,2,2.5,3,4,5,6,8,10,12,24,48],origin='reference single-dose periods of trial2; two reference-period NCA values averaged within person',
            food='Not explicitly stated in this pharmacogenetic report; do not infer fasting.',
            population=dict(healthy=True,n=30,age_range_eligibility=[18,55],bmi_range_eligibility=[18.5,30],pregnant_or_breastfeeding_excluded=True,trial_specific_sex_counts=None,trial_specific_age_weight_means=None,pooled49_participant_demographics_not_substituted_for30=True),
            table='pharmaceutics-13-01573-t003',endpoints=[('AUC_infinity_normalized_by_dose_per_weight',11,2,'kg*ng*h/(mL*mg)','kg*ng*h/(mL*mg)',1),('Cmax_normalized_by_dose_per_weight',11,3,'kg*ng/(mL*mg)','kg*ng/(mL*mg)',1)],
            aggregation='arithmetic_mean_of_individual_dose_per_weight_normalized_NCA_endpoints',papers=['zheng2020'],native_formulation='Solution',
            limitations=['Normalization is endpoint_i/(dose_i/weight_i); do not reconstruct absolute mean endpoints by multiplying by a mean dose/weight.', 'No individual weight linkage or trial2-only mean weight is available; population/reference normalization assumptions must be frozen explicitly.', 'Two reference periods are averaged within each participant;30people do not become60 independent subjects.', 'Current source repository is pregnancy-focused but includes nonpregnant source data; nonpregnant reconstruction must be verified.', 'Food state and between-period washout length are not explicitly reported in the inspected paper.']),
        'vancomycin':dict(pmc='PMC11504268',snapshot='Vancomycin-Model__155fdca3eb25/source/Vancomycin.json',compound='Vancomycin',
            suffix='korean2021_2022_iv1g_2h',trials=[],ethics_id='Hallym2020-12-005',n=12,
            dose_mg=1000,formulation='1g intravenous vancomycin dissolved in200mL normal saline',route='intravenous_infusion',infusion_duration_h=2,dose_times_h=[0],dose_basis='labelled vancomycin1g',
            sampling_h=[2,3,4,6,8],origin='start of2h infusion; no measured0h sample',food='Not reported; IV administration.',
            population=dict(healthy=True,ethnicity='Korean',n=12,sex_counts=dict(male=9,female=3),age_mean_year=35.9,height_mean_cm=170,weight_mean_kg=70.0,bmi_mean=24.0,serum_creatinine_mean_mg_dl=.866,creatinine_clearance_CG_mean_ml_min=117),
            table='antibiotics-13-00993-t003',endpoints=[('Cmax_sampled',2,2,'mg/L','ng/mL',1000),('AUC_last',5,2,'mg*h/L','ng*h/mL',1000),('AUC_infinity',6,2,'mg*h/L','ng*h/mL',1000)],
            aggregation=mean,papers=[],native_formulation=None,
            limitations=['Only5actual sampling times; source authors assumed an unmeasured zero concentration at infusion start for NCA. Preserve this NCA definition, never add it as a measured point.', 'Observed NCA rows only; population-model predicted AUCtau is excluded.', 'Clast/Tlast unit cells in the source table appear swapped; those rows are not endpoints in this packet.', 'Population renal function covariates are means, not individually linked.']),
        'bedaquiline':dict(pmc='PMC10874975',snapshot='Bedaquiline-Model__3351e51d533d/source/Bedaquiline.json',compound='Bedaquiline',
            suffix='south_indian_100mg_fed',trials=[],ethics_id='S2JIndependentEthicsCommitteeHyderabad_no_number_reported',n=6,
            dose_mg=100,formulation='Bedaquiline100mg oral tablet; brand not explicitly reported',route='oral',dose_times_h=[0],dose_basis='labelled bedaquiline dose',
            sampling_h=[0,1,2,3,3.5,4,4.5,5,5.5,6,6.5,7,8,9,10,12,16,24,36,48,72],origin='single100mg oral dose',food='10h overnight fast followed by breakfast; dose30min after serving breakfast.',
            population=dict(healthy=True,sex='male',ethnicity='South Indian',n=6,age_mean_year=None,weight_mean_kg=None,height_mean_cm=None),
            table='t0020',endpoints=[('Cmax',1,1,'ng/mL','ng/mL',1),('AUC_last',3,1,'ng*h/mL','ng*h/mL',1),('AUC_infinity',4,1,'ng*h/mL','ng*h/mL',1)],
            aggregation=mean,papers=['PMC11106169'],native_formulation='Weibull',
            limitations=['Table values are rounded scientific notation; no higher precision fabricated from plot pixels.', 'Age/weight/height and exact brand are not reported, requiring declared reference assumptions.', 'Two incurred-sample-reanalysis points per participant are real sparse measurements selected near peak/tail, not complete independent concentration curves.', 'Parent-M2 source pathway must remain intact.']),
    }
    for drug,compound,dose,table,snapshot,form in [
        ('ethinylestradiol','Ethinylestradiol',.03,'Tab2','Ethinylestradiol-Model__f7aaa8f69206/source/Ethinylestradiol-Model.json','Ethinylestradiol tablet'),
        ('levonorgestrel','Levonorgestrel',.15,'Tab3','Levonorgestrel__bf22ac5e151a/source/Levonorgestrel-Model.json','Microgynon 30'),
    ]:
        cases[drug] = dict(pmc='PMC8416862',snapshot=snapshot,compound=compound,suffix='nct02845219_microgynon_alone_day8',trials=['NCT02845219'],n=25,
            dose_mg=dose,formulation='Microgynon combined tablet:ethinylestradiol0.03mg+levonorgestrel0.15mg',route='oral',dose_times_h=list(range(0,169,24)),dose_basis='labelled active drug within combined oral contraceptive',
            coadministered_drugs=[dict(compound='ethinylestradiol',dose_mg=.03),dict(compound='levonorgestrel',dose_mg=.15)],
            sampling_h=None,observation_origin_h=168,reported_sampling_duration_h=240,origin='last eighth once-daily Microgynon dose in first OC-alone period',
            food='Not explicitly specified in inspected publication/registry; retain unknown.',
            population=dict(healthy=True,postmenopausal=True,sex='female',ethnicity='White',n=25,age_mean_year=62,age_range_year=[50,75],weight_mean_kg=65.5,bmi_mean=24.1,height_mean_cm=None),
            table=table,endpoints=[('AUC_tau_ss',3,1,'pg*h/mL','ng*h/mL',.001),('Cmax_ss',6,1,'pg/mL','ng/mL',.001)],
            aggregation=mean,papers=['PMC7764797'] if drug=='ethinylestradiol' else ['PMC7825189'],native_formulation=form,
            limitations=['OC-alone means combined EE+LNG without semaglutide/SNAC, not single-compound monotherapy.', 'The two drugs share the same25subjects/Trial1 and publication; never counted as independent trials.', 'Detailed nominal sampling grid is not reported; AUCtau is0–24h and dense-Cmax is an explicit approximation until exact grid is available.', 'Postmenopausal62-year population must not silently use young-adult physiology or assume unspecified height.', 'Entire shared publication excluded from development, including its separate Furosemide/Rosuvastatin Trial2.'])
    assert len(cases) == 9
    return cases


def article_path(pmc):
    if pmc == 'PMC8416862':
        return OLD / 'retrieval_round1/sealed_articles/PMC8416862.xml'
    return SOURCE / f'{pmc}.xml'


def paper_path(name):
    if name == 'hanke2018':
        return ROOT / 'experiments/blueprint_design/external_validation/remaining_cohort_20260920/source_audits/model_papers/PMC6202474.xml'
    return PAPERS / f'{name}.xml'


def native_status(spec):
    path = ASSETS / spec['snapshot']
    data = json.loads(path.read_text())
    compound = next(c for c in data['Compounds'] if c['Name'] == spec['compound'])
    adapter = ROOT / 'src/pbpk_kb/blueprints/human_model_snapshot.py'
    tree = ast.parse(adapter.read_text())
    recognized = set()
    names = {'_PROCESSES','_ENZYME_PROCESS_PARAMETERS','_MM_PROCESS_PARAMETERS','_INTERACTION_PROCESS_PARAMETERS','_TRANSPORT_PROCESS_PARAMETERS'}
    for node in tree.body:
        if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id in names for t in node.targets) and isinstance(node.value, ast.Dict):
            recognized.update(k.value for k in node.value.keys if isinstance(k, ast.Constant))
    types = sorted({p['InternalName'] for p in compound.get('Processes', [])})
    formulation_names = [f['Name'] for f in data.get('Formulations', [])]
    if spec['native_formulation'] is not None:
        assert spec['native_formulation'] in formulation_names
    return dict(snapshot=ref(path),compound_name=spec['compound'],source_process_types=types,
                process_types_not_in_inspected_adapter=sorted(set(types)-recognized),adapter_inspected=ref(adapter),
                source_formulation_candidate=spec['native_formulation'],available_formulations=formulation_names,
                source_formulation_matches_target_brand_verified=False,
                additional_source_compounds=[c['Name'] for c in data['Compounds'] if c['Name'] != spec['compound']],
                static_process_allowlist_screen_only=True,actual_native_build_verified=False,
                not_a_complete_semantic_or_execution_qualification=True)


def source_audit(drug, spec, article, inventory):
    ids = {x.get('pub-id-type'): x.text for x in article.findall('.//article-meta/article-id') if x.get('pub-id-type') in {'doi','pmid','pmcid'}}
    title = text(article.find('.//article-meta/title-group/article-title'))
    tokens = [v for v in ids.values() if v] + spec['trials'] + [title]
    docs = {}
    for row in inventory['drug_inventory']:
        if row['compound'] != drug:
            continue
        for model in row['models']:
            p = ROOT / model['snapshot_path']; docs[p] = 'registered_snapshot'
            docs.update({d:'repository_source_document' for d in p.parent.rglob('*.md')})
    p = ASSETS / spec['snapshot'];docs[p] = 'selected_registered_snapshot'
    docs.update({d:'repository_source_document' for d in p.parent.rglob('*.md')})
    for name in spec['papers']:
        if name == 'zheng2020':
            docs[PAPERS/'zheng2020.html']='original_publisher_abstract_all49_references_and_metadata_not_paywalled_body'
            docs[PAPERS/'zheng2020_supplement.docx']='original_model_full_supplement_methods_clinical_source_table_and_references'
        else:
            docs[paper_path(name)]='original_model_full_article_and_references'
    entries=[];overlaps=[];labels=[]
    for p,kind in sorted(docs.items()):
        refs=None
        if p.suffix == '.xml':
            r=ET.parse(p).getroot();content=text(r);refs=len(list(r.iter('ref')))
        elif p.suffix == '.docx':
            d=Document(p);content='\n'.join([x.text for x in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
        elif p.suffix == '.html':
            h=html.fromstring(p.read_bytes())
            for node in h.xpath('//script|//style'):node.drop_tree()
            content=' '.join(h.itertext());refs=len(h.xpath('//li[contains(@class,"c-article-references__item")]'))
        else:
            content=p.read_text()
            if p.suffix=='.json':
                d=json.loads(content);labels += [x.get('Name','') for x in d.get('ObservedData',[])]
        matched=[v for v in tokens if norm(v) in norm(content)]
        record={**ref(p),'type':kind,'target_identity_matches':matched}
        if refs is not None:record['references_checked']=refs
        entries.append(record)
        if matched:overlaps.append(record)
    if overlaps:
        save(OUT/'source_audits'/f'{drug}_overlap.json',dict(compound=drug,clinical_source_independence_supported=False,overlaps=overlaps))
        raise ValueError('Clinical source overlap detected for '+drug)
    return dict(schema_version='pbpk-clinical-source-independence-audit-v1',compound=drug,
                target_identifiers=ids,target_title=title,target_trial_ids=spec['trials'],
                selected_model=ref(ASSETS/spec['snapshot']),audit_documents=entries,
                registered_historical_study_labels=sorted(set(labels)),known_clinical_overlap=[],
                clinical_source_independence_supported_at_declared_scope=True,
                source_qualification_scope='All registered same-drug snapshots and repository source documents; linked original-model papers or full clinical-source supplement plus publisher reference list. Separate clinical cohort/source identity, not generic physiological-literature exclusivity.',
                old_and_new_trial_aliases_checked=True,
                no_guarantee_of_undocumented_individual_reparticipation=True,
                no_target_values_used_in_native_model_development=True)


def scalar(token):
    # Source scientific notation appears as 1.16x10<sup>3</sup> in XML.
    token=token.replace(',','').strip()
    m=re.match(r'([0-9]+(?:\.[0-9]+)?)\s*[x×]\s*10\s*([+-]?\d+)',token)
    if m:return float(m[1])*10**int(m[2])
    m=re.match(r'([0-9]+(?:\.[0-9]+)?)',token)
    if not m:raise ValueError('Cannot parse source scalar')
    return float(m[1])


def main():
    os.sched_setaffinity(0,sorted(os.sched_getaffinity(0))[:2])
    if (OUT/'curation_receipt.json').exists():raise SystemExit('Completed curation exists; never overwrite frozen inputs')
    inventory=json.loads((OLD/'cohort_inventory.json').read_text())
    old_summary=json.loads((OLD/'acquisition_summary.json').read_text())
    existing=set(old_summary['curation_coverage']['drug_names'])
    assert not existing.intersection(specifications())
    records=[];validation=[];source_errors=[];deny=[]
    for drug,spec in specifications().items():
        article_file=article_path(spec['pmc']);article=ET.parse(article_file).getroot()
        audit=source_audit(drug,spec,article,inventory)
        audit_ref=save(OUT/'source_audits'/f'{drug}.json',audit)
        case_id=f"{drug}_{spec['pmc'].lower()}_{spec['suffix']}"
        if spec.get('docx'):
            source_table_file=OUT/'sealed/supplements'/spec['docx'];table=Document(source_table_file).tables[spec['table']]
            cells=[[c.text for c in row.cells] for row in table.rows]
            original_table={'docx':ref(source_table_file),'table_index':spec['table'],'rows':cells}
        else:
            source_table_file=article_file;table=next(t for t in article.iter('table-wrap') if t.get('id')==spec['table'])
            cells=[[text(c) for c in row if c.tag in {'td','th'}] for row in table.iter('tr')]
            original_table={'source_article':ref(article_file),'table_id':spec['table'],'source_table_xml':ET.tostring(table,encoding='unicode')}
        endpoints=[];contracts=[]
        for name,row,col,unit,comp_unit,mult in spec['endpoints']:
            token=cells[row][col];value=scalar(token);assert value>0
            reason=spec.get('disabled_endpoints',{}).get(name)
            ep=dict(endpoint=name,value=value,source_cell=token,row_0based=row,column_0based=col,table_id=spec['table'],
                    source_unit=unit,comparable_unit=comp_unit,conversion_multiplier=mult,comparable_value=value*mult,
                    n_evaluable=spec['n'],aggregation=spec['aggregation'],evaluable_from_source=reason is None,source_quality_blocker=reason)
            endpoints.append(ep)
            contracts.append({k:v for k,v in ep.items() if k not in {'value','source_cell','comparable_value'}})
            if reason:source_errors.append(dict(compound=drug,endpoint=name,reason=reason,original_value_preserved_sealed=True))
        if drug=='omeprazole':
            by={x['endpoint']:x['value'] for x in endpoints}
            assert by['AUC_infinity']<by['AUC_last'], 'Recheck source-quality flag if the primary source changes'
        sealed=save(OUT/'sealed/endpoints'/f'{case_id}.json',dict(schema_version='pbpk-independent-endpoints-sealed-v1',case_id=case_id,compound=drug,source_identifiers=audit['target_identifiers'],source_article=ref(article_file),source_table=original_table,endpoints=endpoints),True)
        native=native_status(spec)
        blockers=[]
        if native['process_types_not_in_inspected_adapter']:blockers.append('Unsupported native process types:'+','.join(native['process_types_not_in_inspected_adapter']))
        if drug=='quetiapine':blockers.append('Normalized individual endpoint aggregation needs linked weight information or an explicitly frozen reference/population approximation; absolute AUC/Cmax cannot be recovered from this summary.')
        if drug in {'ethinylestradiol','levonorgestrel'}:blockers.append('Exact sample grid absent; register dense-Cmax approximation and combined-hormone/binding assumptions before prediction.')
        protocol={
            'schema_version':'pbpk-independent-clinical-design-v2','case_id':case_id,'compound':drug,
            'source_identifiers':audit['target_identifiers'],'publication_cluster_id':spec['pmc'],
            'cohort_id': 'NCT02845219:OC_alone_period1' if drug in {'ethinylestradiol','levonorgestrel'} else f"{spec['pmc']}:{spec['suffix']}",
            'trial_registration_ids_at_publication_level':spec['trials'],'ethics_id':spec.get('ethics_id'),
            'dose':dict(value=spec['dose_mg'],unit='mg',basis=spec['dose_basis'],route=spec['route'],formulation=spec['formulation'],administration_times_h=spec['dose_times_h'],infusion_duration_h=spec.get('infusion_duration_h'),coadministered_drugs=spec.get('coadministered_drugs',[])),
            'food':spec['food'],'population':spec['population'],
            'sampling':dict(times_h_from_observation_origin=spec['sampling_h'],origin=spec['origin'],observation_origin_h_after_first_dose=spec.get('observation_origin_h',0),predose_interval_h=spec.get('predose_interval_h'),reported_sampling_duration_h=spec.get('reported_sampling_duration_h'),missing_nominal_grid_not_fabricated=spec['sampling_h'] is None),
            'matrix':'plasma','analyte':'total parent drug; total racemate for omeprazole','endpoint_contracts':contracts,
            'aggregation_preselected_before_predictions':True,'auc_definition':{'AUC_last':'integral to last individual quantifiable sample; not automatically final scheduled sample','AUC_infinity':'individual NCA total with extrapolated terminal tail','AUC_tau_ss':'last-dose0–24h interval for eight-day contraceptive schedule','source_specific_normalization':'quetiapine endpoint_i/(dose_i/weight_i), then within-person reference-period average and across-person mean' if drug=='quetiapine' else None},
            'native_binding':native,'source_independence_audit':audit_ref,'sealed_observation_reference':sealed,
            'curator_only_source_article':ref(article_file),'evidence_exclusion':dict(entire_publication=spec['pmc'],doi=audit['target_identifiers'].get('doi'),pmid=audit['target_identifiers'].get('pmid'),all_trial_arms_and_derived_results=True,trial_ids=spec['trials']),
            'source_metadata_and_endpoints_ready':True,'independent_source_qualification_supported':True,'native_execution_verified':False,'run_ready':False,'actual_evaluation_completed':False,'strict_main_evaluated':False,
            'execution_requirements':blockers+['Verify source-native reconstruction and formulation/population/food matching without outcomes.','Freeze all assumptions and selected predictions before opening sealed endpoints.','Use aggregation-specific scoring; do not compare individual-normalized or geometric quantities against incompatible means.'],
            'limitations':spec['limitations'],'observed_values_exposed_to_method_development':False,
        }
        if drug=='vancomycin':protocol['source_nca_baseline_handling']={'measured_baseline_available':False,'source_assumed_concentration_at_infusion_start':0,'assumed_value_is_not_a_measured_observation':True,'comparison_requires_registering_same_sparse_NCA_rule_or_explicit_approximation':True}
        pref=save(OUT/'protocols'/f'{case_id}.json',protocol)
        records.append(dict(compound=drug,case_id=case_id,pmcid=spec['pmc'],cohort_id=protocol['cohort_id'],protocol=pref,source_audit=audit_ref,sealed_endpoints=sealed,endpoint_count=len(endpoints),source_evaluable_endpoint_count=sum(e['evaluable_from_source'] for e in endpoints),native_verified=False,run_ready=False,evaluated=False,execution_requirements=blockers))
        validation.append(dict(case_id=case_id,exact_source_cell_mapping=True,units_registered=True,source_aggregation_preserved=True,source_model_exists=True,clinical_source_audit_passed=True,source_errors_not_corrected_or_hidden=True))
        deny.append(protocol['evidence_exclusion'])
    save(OUT/'evidence_denylist.json',{'exclude_from_all_development_retrieval_prompts_training_and_selection':deny,'publication_level_clustering_required':True})
    save(OUT/'static_validation.json',dict(status='passed',cases=validation,solver_calls=0,gpu_used=False))
    save(OUT/'source_quality_audit.json',dict(disabled_endpoints=source_errors,values_preserved_not_corrected=True))
    receipt=dict(schema_version='pbpk-expanded-clinical-curation-receipt-v1',status='nine_new_drugs_source_metadata_and_sealed_endpoints_prepared',records=records,
                 new_unique_drugs=len(records),new_protocol_records=len(records),new_target_publications_in_this_batch=len({r['pmcid'] for r in records}),
                 target_cohorts_in_this_batch=len({r['cohort_id'] for r in records}),existing_candidate_drugs_before=21,total_candidate_drugs_after=30,
                 unique_new_publications_beyond_old20=7,shared_previously_curated_publication='PMC8416862 with new Trial1 EE/LNG cohort, distinct from prior Trial2 Furosemide',
                 sealed_endpoints=sum(r['endpoint_count'] for r in records),source_evaluable_endpoints=sum(r['source_evaluable_endpoint_count'] for r in records),
                 native_verified_cases=0,actual_evaluated_cases=0,full_independent_numeric_curves_added=0,
                 prepared_protocol_records_are_not_run_ready_count=True,download_footprint_bytes=cohort_bytes(),maximum_cumulative_cohort_bytes=100*1024**2,
                 no_model_prediction_or_score_files_read=True,solver_calls=0,gpu_used=False,script=ref(Path(__file__)))
    save(OUT/'curation_receipt.json',receipt)
    print(json.dumps({k:receipt[k] for k in ['status','new_unique_drugs','new_protocol_records','sealed_endpoints','source_evaluable_endpoints','total_candidate_drugs_after','native_verified_cases','actual_evaluated_cases']},ensure_ascii=False))


if __name__=='__main__':main()
